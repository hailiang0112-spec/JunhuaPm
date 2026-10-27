"""计划 Excel / CSV 导入导出（REQ-F03-04，API-F03-03）与原型 v1.20 JSON 备份导入（F18，API-F18-01）。

Excel：XlsxWriter 写、openpyxl 读；5 张工作表（施工阶段、工作事项、补充事项、填写说明、取值），数值列为数字格式、
日期列为文本格式（AC-F03-07）。按表头文字识别列，逐行校验并给出工作表与行号（AC-F03-08）；导入只新增工作项或更新
未签认的，已签认的列出「已签认，需走变更单」（AC-F03-09）。
原型 JSON：只允许「筹备」项目；每条标注 prototype_import；不生成签认、核对或放行记录；只有元数据的文件标记缺失；
无效文件拒绝且整个导入事务回滚（AC-F18-01 ~ AC-F18-03）。
"""
import base64
import csv
import datetime
import io
import json
import re
from pathlib import Path

from django.db import models, transaction
from django.http import HttpResponse
from django.urls import path
from rest_framework.response import Response

from core.api import structure as st
from core.api.base import ProjectView, d, iso
from core.errors import ApiError, fail
from core.models import (AcceptancePhoto, ArchiveFile, ArchiveFolder, Attendance, Baseline, Floor, ImportBatch, Issue,
                         Notice, Phase, PlanChangeLog, Room, StoredFile, Supplement, WorkItem, WorkItemAttachment)
from core.services import audit, ledger, storage
from core.services.projects import _date_in_year
from core.services.shared import issue_event, next_issue_seq

LOCKED_BASELINE = ("待签认", "已签认", "待重签")
SIGNED_NOTE = "已签认，需走变更单"

# 列定义：表头、键、宽度、类型（d=文本日期 n=数字 w=长文本）——模板、导出、解析共用
PHASE_COLS = [("序号", "idx", 6, "n"), ("阶段名称", "name", 14, ""), ("计划开始日期", "plannedStart", 15, "d"),
              ("计划结束日期", "plannedEnd", 15, "d"), ("施工班组 / 负责人", "teamLead", 26, ""),
              ("阶段权重(%)", "weight", 12, "n"), ("施工范围说明", "scopeNote", 46, "w")]
TASK_COLS = [("序号", "idx", 6, "n"), ("施工阶段", "phase", 12, ""), ("工作事项名称", "name", 34, ""),
             ("工作说明", "description", 44, "w"), ("验收标准", "acceptanceCriteria", 44, "w"),
             ("所属楼层", "floor", 13, ""), ("所属房间 / 区域", "room", 18, ""), ("计划完成日期", "plannedFinish", 15, "d"),
             ("实际完成日期", "actualFinish", 15, "d"), ("负责人", "ownerName", 12, ""), ("权重(1-5)", "weight", 10, "n"),
             ("完成度(%)", "progress", 11, "n")]
EXTRA_COLS = [("序号", "idx", 6, "n"), ("事项标题", "title", 28, "w"), ("说明", "note", 44, "w"),
              ("提出日期", "raisedOn", 14, "d"), ("约定完成日期", "dueOn", 14, "d"), ("责任人", "ownerName", 12, ""),
              ("状态", "status", 12, "")]
# 表头别名（按表头文字识别列，列顺序可调）
ALIASES = {
    "phase_sheet": {"name": ["阶段名称", "阶段", "名称"], "plannedStart": ["计划开始日期", "开始日期", "计划开始"],
                    "plannedEnd": ["计划结束日期", "结束日期", "计划结束"],
                    "teamLead": ["施工班组/负责人", "施工班组", "班组/负责人", "班组负责人", "班组"],
                    "weight": ["阶段权重(%)", "阶段权重", "权重(%)", "权重"], "scopeNote": ["施工范围说明", "施工范围", "范围"]},
    "task_sheet": {"phase": ["施工阶段", "阶段", "所属阶段", "工段"], "name": ["工作事项名称", "工作事项", "事项名称", "工作项名称"],
                   "description": ["工作说明", "说明"], "acceptanceCriteria": ["验收标准", "标准"],
                   "floor": ["所属楼层", "楼层"], "room": ["所属房间/区域", "所属房间", "房间/区域", "房间"],
                   "plannedFinish": ["计划完成日期", "计划日期", "计划"], "actualFinish": ["实际完成日期", "实际日期", "实际"],
                   "ownerName": ["负责人", "责任人"], "weight": ["权重(1-5)", "权重"],
                   "progress": ["完成度(%)", "完成度", "完成率(%)", "完成率"]},
    "extra_sheet": {"title": ["事项标题", "标题"], "note": ["说明", "描述"], "raisedOn": ["提出日期", "日期"],
                    "dueOn": ["约定完成日期", "约定日期", "截止日期"], "ownerName": ["责任人", "负责人"], "status": ["状态"]},
}
SUPPLEMENT_STATUS = ["待处理", "处理中", "已完成"]
TASK_LABEL = {"description": "工作说明", "acceptanceCriteria": "验收标准", "floor": "楼层", "room": "房间",
              "plannedFinish": "计划日期", "actualFinish": "实际完成日期", "ownerName": "负责人", "weight": "权重",
              "progress": "完成度"}


def norm(s):
    return re.sub(r"[\s　]", "", str("" if s is None else s)).replace("（", "(").replace("）", ")") \
        .replace("：", ":").replace("／", "/").replace("－", "-").replace("—", "-")


def phase_key(s):
    return re.sub(r"(阶段|工段|工序)$", "", norm(s))


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_cell_date(v, year_hint=None, span=None):
    """返回 (date | None, 是否有值)。支持日期单元格、YYYY-MM-DD、YYYY/M/D、YYYY年M月D日、只有月日（按工段年份补全）。"""
    if isinstance(v, datetime.datetime):
        return v.date(), True
    if isinstance(v, datetime.date):
        return v, True
    s = cell_text(v)
    if not s:
        return None, False
    s2 = s.replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s2)
    try:
        if m:
            return datetime.date(int(m[1]), int(m[2]), int(m[3])), True
        m = re.fullmatch(r"(\d{1,2})-(\d{1,2})", s2)
        if m:
            datetime.date(2000, int(m[1]), int(m[2]))  # 校验月日合法（含 2 月 29 日）
            if span:
                return _date_in_year(f"{int(m[1])}-{int(m[2])}", span[0], span[1]), True
            return datetime.date(year_hint or datetime.date.today().year, int(m[1]), int(m[2])), True
    except ValueError:
        return None, True
    return None, True


def parse_int(v):
    s = cell_text(v)
    if s == "":
        return None, False
    try:
        f = float(s)
    except ValueError:
        return "nan", True
    return (int(f) if f.is_integer() else f), True


# ---------------------------------------------------------------- 写 xlsx

def build_workbook(project, mode):
    import xlsxwriter
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    head = wb.add_format({"bold": True, "bg_color": "#EEF2F7", "border": 1, "text_wrap": True, "valign": "vcenter"})
    text = wb.add_format({"num_format": "@"})
    wrap = wb.add_format({"text_wrap": True, "valign": "top"})
    num = wb.add_format({"num_format": "0"})
    phases = list(Phase.objects.filter(project=project, is_disabled=False).order_by("order_no", "id"))
    floors = list(Floor.objects.filter(project=project, is_disabled=False).order_by("order_no", "id"))
    rooms = list(Room.objects.filter(project=project, is_disabled=False, floor__is_disabled=False)
                 .select_related("floor").order_by("floor__order_no", "order_no", "id"))

    def sheet(name, cols, rows):
        ws = wb.add_worksheet(name)
        for i, (h, _, w, t) in enumerate(cols):
            ws.set_column(i, i, w, text if t == "d" else wrap if t == "w" else None)
            ws.write_string(0, i, h, head)
        for r, row in enumerate(rows, start=1):
            for i, (_, key, _, t) in enumerate(cols):
                v = row.get(key)
                if v is None or v == "":
                    if t == "d":
                        ws.write_blank(r, i, None, text)
                    continue
                if t == "n":
                    ws.write_number(r, i, float(v), num)
                elif t == "d":
                    ws.write_string(r, i, d(v) if isinstance(v, datetime.date) else str(v), text)
                else:
                    ws.write_string(r, i, str(v), wrap if t == "w" else None)
        ws.freeze_panes(1, 0)
        ws.autofilter(0, 0, max(len(rows), 1), len(cols) - 1)
        return ws

    prow = [{"idx": i + 1, "name": p.name, "plannedStart": p.planned_start, "plannedEnd": p.planned_end,
             "teamLead": p.team_lead, "weight": p.weight, "scopeNote": p.scope_note} for i, p in enumerate(phases)]
    trow, xrow = [], []
    if mode == "export":
        fl = {f.id: f.name for f in Floor.objects.filter(project=project)}
        rm = {r.id: r.name for r in Room.objects.filter(project=project)}
        ph = {p.id: p.name for p in phases}
        qs = WorkItem.objects.filter(project=project, is_disabled=False, phase_id__in=list(ph)).order_by("phase__order_no", "id")
        for i, w in enumerate(qs):
            trow.append({"idx": i + 1, "phase": ph[w.phase_id], "name": w.name, "description": w.description,
                         "acceptanceCriteria": w.acceptance_criteria, "floor": fl.get(w.floor_id, "全楼通用"),
                         "room": rm.get(w.room_id, "本层通用" if w.floor_id else ""), "plannedFinish": w.planned_finish,
                         "actualFinish": w.actual_finish, "ownerName": w.owner_name, "weight": w.weight,
                         "progress": w.progress})
        for i, s in enumerate(Supplement.objects.filter(project=project).order_by("id")):
            xrow.append({"idx": i + 1, "title": s.title, "note": s.note, "raisedOn": s.raised_on, "dueOn": s.due_on,
                         "ownerName": s.owner_name, "status": s.status})
    sheet("施工阶段", PHASE_COLS, prow)
    sheet("工作事项", TASK_COLS, trow)
    sheet("补充事项", EXTRA_COLS, xrow)
    ex = wb.add_worksheet("填写说明")
    ex.set_column(0, 0, 18)
    ex.set_column(1, 1, 14)
    ex.set_column(2, 2, 86, wrap)
    explain = [
        ["君华装修项目管理平台 · " + ("当前项目计划（导出件）" if mode == "export" else "项目计划导入模板"), "", ""],
        ["项目", project.name, ""],
        ["", "", ""],
        ["工作表", "是否必填", "说明"],
        ["施工阶段", "选填", "工段计划字段（起止日期、班组、权重、施工范围）。阶段名称须与「取值」表一致；有变化的工段导入后自动写入计划变更留档。"],
        ["工作事项", "必填", "每行一条工作事项。按「施工阶段 + 工作事项名称」匹配已有工作项：匹配到且未签认的按本表更新；匹配不到的新增；已签认的不做任何修改，结果中列为「已签认，需走变更单」。"],
        ["补充事项", "选填", "只新增标题不重复的补充事项；已有的不修改。"],
        ["取值", "参考", "工段、楼层与房间对照、补充事项状态的合法取值。"],
        ["", "", ""],
        ["日期格式", "YYYY-MM-DD", "例 2026-09-19。日期列已设为文本格式，不会被 Excel 转成数字；只写 09-19 时按所属工段年份补全。"],
        ["所属楼层", "可留空", "留空或填「全楼通用」= 跨楼层通项。"],
        ["所属房间", "可留空", "留空或填「本层通用」= 整层通用。房间必须属于所填楼层，否则校验报错。"],
        ["权重(1-5)", "可留空", "1–5 的整数；新增时留空按 2 处理。"],
        ["完成度(%)", "可留空", "0–100 的整数；新增时留空按「实际完成日期有无」推 100 或 0。"],
        ["阶段权重(%)", "可留空", "1–100 的整数。"],
        ["校验", "", "按表头文字识别列，请勿修改表头；任一行有错误时整份文件不能导入，错误清单列出工作表与行号。"],
    ]
    for r, row in enumerate(explain):
        for c, v in enumerate(row):
            ex.write_string(r, c, v, head if r in (0, 3) else wrap if c == 2 else None)
    vs = wb.add_worksheet("取值")
    for i, h in enumerate(["工段", "楼层", "房间（楼层 / 房间）", "补充事项状态"]):
        vs.set_column(i, i, 22)
        vs.write_string(0, i, h, head)
    for r, p in enumerate(phases, start=1):
        vs.write_string(r, 0, p.name)
    for r, f in enumerate(floors + [None], start=1):
        vs.write_string(r, 1, f.name if f else "全楼通用")
    for r, rm_ in enumerate(rooms, start=1):
        vs.write_string(r, 2, f"{rm_.floor.name} / {rm_.name}")
    for r, s in enumerate(SUPPLEMENT_STATUS, start=1):
        vs.write_string(r, 3, s)
    wb.close()
    return buf.getvalue()


def xlsx_response(data, filename):
    from urllib.parse import quote
    resp = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return resp


class PlanTemplate(ProjectView):
    def get(self, request, project_id):
        self.require("plan.import", "只有项目经理可以导入导出计划")
        return xlsx_response(build_workbook(self.project, "template"), f"{self.project.name}-空白计划模板.xlsx")


class PlanExport(ProjectView):
    def get(self, request, project_id):
        self.require("plan.import", "只有项目经理可以导入导出计划")
        return xlsx_response(build_workbook(self.project, "export"),
                             f"{self.project.name}-当前计划-{d(st.today())}.xlsx")


# ---------------------------------------------------------------- 读与校验

def read_workbook(name, data):
    """返回 {工作表名: [行(list)]}。xlsx 用 openpyxl 只读；csv 自动识别 UTF-8 / GBK。"""
    ext = Path(name).suffix.lower()
    if ext == ".csv":
        for enc in ("utf-8-sig", "gbk"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("csv")
        return {"CSV": [list(r) for r in csv.reader(io.StringIO(text))]}
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    wb.close()
    return out


def classify(sheets):
    out = {}
    for nm, rows in sheets.items():
        if not rows:
            continue
        head = [norm(c) for c in rows[0]]
        if "阶段名称" in head:
            out.setdefault("phase_sheet", (nm, rows))
        elif "工作事项名称" in head or "工作项名称" in head:
            out.setdefault("task_sheet", (nm, rows))
        elif "事项标题" in head:
            out.setdefault("extra_sheet", (nm, rows))
    return out


def find_cols(head, kind):
    h = [norm(c) for c in head]
    cols = {}
    for key, names in ALIASES[kind].items():
        for n in names:
            if norm(n) in h:
                cols[key] = (h.index(norm(n)), head[h.index(norm(n))])
                break
    return cols


def blank(row):
    return all(cell_text(c) == "" for c in row)


class Validator:
    def __init__(self, project):
        self.project = project
        self.phases = list(Phase.objects.filter(project=project, is_disabled=False).order_by("order_no", "id"))
        self.floors = list(Floor.objects.filter(project=project, is_disabled=False))
        self.rooms = list(Room.objects.filter(project=project, is_disabled=False, floor__is_disabled=False))
        # 已停用的楼层与房间：只允许已有工作项沿用原位置（本系统导出的文件可原样导回），新增工作项不能放进去
        self.disabled_floors = list(Floor.objects.filter(project=project, is_disabled=True))
        self.disabled_rooms = list(Room.objects.filter(project=project).filter(models.Q(is_disabled=True) | models.Q(floor__is_disabled=True)))
        self.existing = {(w["phase_id"], w["name"]): (w["floor_id"], w["room_id"])
                         for w in WorkItem.objects.filter(project=project).values("phase_id", "name", "floor_id", "room_id")}
        self.errors = []

    def err(self, sheet, row, col, msg):
        self.errors.append({"sheet": sheet, "row": row, "column": str(col or ""), "message": msg})

    def phase_of(self, s):
        k = phase_key(s)
        return next((p for p in self.phases if phase_key(p.name) == k), None) if k else None

    def phase_names(self):
        return "、".join(p.name for p in self.phases)

    def run(self, sheets):
        cls = classify(sheets)
        parsed = {"phases": [], "tasks": [], "extras": [], "sheets": len(sheets),
                  "taskRows": 0, "hasPhaseSheet": "phase_sheet" in cls}
        if "task_sheet" not in cls:
            self.err("工作事项", 1, "工作事项名称", "未找到「工作事项」表（表头需含「工作事项名称」列）")
        if "phase_sheet" in cls:
            self.phase_sheet(*cls["phase_sheet"], parsed)
        if "task_sheet" in cls:
            self.task_sheet(*cls["task_sheet"], parsed)
        if "extra_sheet" in cls:
            self.extra_sheet(*cls["extra_sheet"], parsed)
        return parsed

    def phase_sheet(self, nm, rows, parsed):
        sheet = "施工阶段"
        cols = find_cols(rows[0], "phase_sheet")
        if "name" not in cols:
            self.err(sheet, 1, "阶段名称", "缺少「阶段名称」列")
            return
        seen = {}
        for i, row in enumerate(rows[1:], start=2):
            if blank(row):
                continue

            def get(k):
                return row[cols[k][0]] if k in cols and cols[k][0] < len(row) else None
            name = cell_text(get("name"))
            if not name:
                continue
            p = self.phase_of(name)
            if p is None:
                self.err(sheet, i, cols["name"][1], f"阶段名「{name}」不可识别，可选：{self.phase_names()}")
                continue
            if p.id in seen:
                self.err(sheet, i, cols["name"][1], f"阶段「{name}」与第 {seen[p.id]} 行重复")
                continue
            seen[p.id] = i
            rec = {"row": i, "phaseId": p.id}
            for key in ("plannedStart", "plannedEnd"):
                if key in cols:
                    v, has = parse_cell_date(get(key), p.planned_start.year)
                    if has and v is None:
                        self.err(sheet, i, cols[key][1], f"日期「{cell_text(get(key))}」无法识别")
                    elif v:
                        rec[key] = d(v)
            if rec.get("plannedStart") and rec.get("plannedEnd") and rec["plannedStart"] > rec["plannedEnd"]:
                self.err(sheet, i, cols["plannedStart"][1], "开始日期不能晚于结束日期")
            if "teamLead" in cols:
                rec["teamLead"] = cell_text(get("teamLead"))[:60]
            if "scopeNote" in cols:
                rec["scopeNote"] = cell_text(get("scopeNote"))
            if "weight" in cols:
                w, has = parse_int(get("weight"))
                if has and (not isinstance(w, int) or not 1 <= w <= 100):
                    self.err(sheet, i, cols["weight"][1], f"阶段权重「{cell_text(get('weight'))}」越界，应为 1–100 的整数")
                elif has:
                    rec["weight"] = w
            parsed["phases"].append(rec)

    def task_sheet(self, nm, rows, parsed):
        sheet = "工作事项"
        cols = find_cols(rows[0], "task_sheet")
        if "name" not in cols:
            self.err(sheet, 1, "工作事项名称", "缺少「工作事项名称」列")
            return
        seen = {}
        for i, row in enumerate(rows[1:], start=2):
            if blank(row):
                continue
            parsed["taskRows"] += 1

            def get(k):
                return row[cols[k][0]] if k in cols and cols[k][0] < len(row) else None

            def col(k):
                return cols[k][1] if k in cols else TASK_LABEL.get(k, k)
            name = cell_text(get("name"))
            if not name:
                self.err(sheet, i, col("name"), "工作事项名称为空")
                continue
            pv = cell_text(get("phase"))
            p = self.phase_of(pv) if "phase" in cols else (self.phases[0] if self.phases else None)
            if p is None:
                self.err(sheet, i, col("phase"), f"阶段名「{pv}」不可识别，可选：{self.phase_names()}")
                continue
            rec = {"row": i, "phaseId": p.id, "name": name[:120]}
            if (p.id, name) in seen:
                self.err(sheet, i, col("name"), f"工作事项「{name}」与第 {seen[(p.id, name)]} 行重复")
                continue
            seen[(p.id, name)] = i
            # 楼层与房间（AC-F03-04 同口径）
            fv, rv = cell_text(get("floor")), cell_text(get("room"))
            if "/" in rv and len(rv.split("/")) == 2:
                a, b = rv.split("/")
                fv, rv = (fv or a.strip()), b.strip()
            floor = None
            ok = True
            kept_floor, kept_room = self.existing.get((p.id, name), (None, None))
            if fv and not re.fullmatch(r"(全楼通用|通用|全部|跨楼层|全楼)", norm(fv)):
                floor = next((f for f in self.floors if norm(f.name) == norm(fv)), None) or next(
                    (f for f in self.disabled_floors if norm(f.name) == norm(fv) and f.id == kept_floor), None)
                if floor is None:
                    self.err(sheet, i, col("floor"), f"楼层「{fv}」不可识别")
                    ok = False
            room = None
            if ok and rv and not re.fullmatch(r"(本层通用|通用|整层|全层)", norm(rv)):
                hits = [r for r in self.rooms if norm(r.name) == norm(rv)] or [
                    r for r in self.disabled_rooms if norm(r.name) == norm(rv) and r.id == kept_room]
                if not hits:
                    self.err(sheet, i, col("room"), f"房间「{rv}」不可识别")
                    ok = False
                elif floor is None:
                    if len(hits) > 1:
                        self.err(sheet, i, col("room"), f"房间「{rv}」在多个楼层重名，请同时填写所属楼层")
                        ok = False
                    elif fv:  # 填了「全楼通用」又填了具体房间
                        self.err(sheet, i, col("room"), f"房间「{rv}」不属于「全楼通用」，房间与楼层不匹配")
                        ok = False
                    else:
                        room = hits[0]
                        floor = next(f for f in self.floors + self.disabled_floors if f.id == room.floor_id)
                else:
                    room = next((r for r in hits if r.floor_id == floor.id), None)
                    if room is None:
                        self.err(sheet, i, col("room"), f"房间「{rv}」不属于楼层「{floor.name}」，房间与楼层不匹配")
                        ok = False
            if ok and ("floor" in cols or "room" in cols):
                rec["floorId"] = floor.id if floor else None
                rec["roomId"] = room.id if room else None
            for key in ("plannedFinish", "actualFinish"):
                if key in cols:
                    v, has = parse_cell_date(get(key), span=(p.planned_start, p.planned_end))
                    if has and v is None:
                        self.err(sheet, i, col(key), f"日期「{cell_text(get(key))}」无法识别")
                    elif v:
                        rec[key] = d(v)
            w, has = parse_int(get("weight"))
            if has and (not isinstance(w, int) or not 1 <= w <= 5):
                self.err(sheet, i, col("weight"), f"权重「{cell_text(get('weight'))}」越界，应为 1–5 的整数")
            elif has:
                rec["weight"] = w
            pc, has = parse_int(get("progress"))
            if has and (not isinstance(pc, int) or not 0 <= pc <= 100):
                self.err(sheet, i, col("progress"), f"完成度「{cell_text(get('progress'))}」越界，应为 0–100 的整数")
            elif has:
                rec["progress"] = pc
            for key in ("description", "acceptanceCriteria", "ownerName"):
                if key in cols and cell_text(get(key)):
                    rec[key] = cell_text(get(key))[:40] if key == "ownerName" else cell_text(get(key))
            parsed["tasks"].append(rec)
        if not parsed["tasks"] and not any(e["sheet"] == sheet for e in self.errors):
            self.err(sheet, 2, "工作事项名称", "「工作事项」表没有任何有效数据行")

    def extra_sheet(self, nm, rows, parsed):
        sheet = "补充事项"
        cols = find_cols(rows[0], "extra_sheet")
        if "title" not in cols:
            return
        for i, row in enumerate(rows[1:], start=2):
            if blank(row):
                continue

            def get(k):
                return row[cols[k][0]] if k in cols and cols[k][0] < len(row) else None
            title = cell_text(get("title"))
            if not title:
                self.err(sheet, i, cols["title"][1], "事项标题为空")
                continue
            rec = {"row": i, "title": title[:100], "note": cell_text(get("note")),
                   "ownerName": cell_text(get("ownerName"))[:40] or "待指派",
                   "status": cell_text(get("status")) or "待处理"}
            if rec["status"] not in SUPPLEMENT_STATUS:
                self.err(sheet, i, cols["status"][1], f"状态「{rec['status']}」不合法，应为 待处理 / 处理中 / 已完成")
                continue
            for key in ("raisedOn", "dueOn"):
                if key in cols:
                    v, has = parse_cell_date(get(key), st.today().year)
                    if has and v is None:
                        self.err(sheet, i, cols[key][1], f"日期「{cell_text(get(key))}」无法识别")
                    elif v:
                        rec[key] = d(v)
            parsed["extras"].append(rec)


# ---------------------------------------------------------------- 预览与执行

def _fmt(key, v, names):
    if v is None or v == "":
        return "（空）"
    if key == "floorId":
        return names["floor"].get(v, "")
    if key == "roomId":
        return names["room"].get(v, "")
    if key in ("plannedFinish", "actualFinish"):
        return str(v)[5:] if len(str(v)) == 10 else str(v)
    if key in ("description", "acceptanceCriteria"):
        return "（改写）"
    return str(v)


def classify_tasks(project, parsed):
    """对比当前库：新增 / 更新未签认 / 已签认不改 / 与当前一致。"""
    names = {"floor": {f.id: f.name for f in Floor.objects.filter(project=project)},
             "room": {r.id: r.name for r in Room.objects.filter(project=project)}}
    names["floor"][None] = "全楼通用"
    names["room"][None] = "本层通用"
    existing = {}
    for w in WorkItem.objects.filter(project=project, is_disabled=False).select_related("baseline").order_by("id"):
        existing.setdefault((w.phase_id, w.name), w)
    out = {"add": [], "update": [], "signed": [], "same": []}
    for t in parsed["tasks"]:
        w = existing.get((t["phaseId"], t["name"]))
        if w is None:
            out["add"].append({"row": t["row"], "name": t["name"]})
            continue
        cur = {"floorId": w.floor_id, "roomId": w.room_id, "plannedFinish": d(w.planned_finish),
               "actualFinish": d(w.actual_finish), "ownerName": w.owner_name, "weight": w.weight,
               "progress": w.progress, "description": w.description, "acceptanceCriteria": w.acceptance_criteria}
        diffs = []
        for key, label in (("floorId", "楼层"), ("roomId", "房间"), ("plannedFinish", "计划日期"),
                           ("actualFinish", "实际完成日期"), ("ownerName", "负责人"), ("weight", "权重"),
                           ("progress", "完成度"), ("description", "工作说明"), ("acceptanceCriteria", "验收标准")):
            if key in t and t[key] != cur[key]:
                if key in ("description", "acceptanceCriteria"):
                    diffs.append(f"{label}（改写）")
                else:
                    diffs.append(f"{label} {_fmt(key, cur[key], names)} → {_fmt(key, t[key], names)}")
        entry = {"row": t["row"], "name": t["name"], "workItemId": w.id, "changes": diffs}
        status = w.baseline.status if hasattr(w, "baseline") else "草稿"
        if not diffs:
            out["same"].append(entry)
        elif status in LOCKED_BASELINE:
            out["signed"].append({**entry, "baselineStatus": status, "note": SIGNED_NOTE})
        else:
            out["update"].append(entry)
    phases = {p.id: p for p in Phase.objects.filter(project=project)}
    out["phaseChanges"] = []
    for rec in parsed["phases"]:
        p = phases[rec["phaseId"]]
        cur = {"plannedStart": d(p.planned_start), "plannedEnd": d(p.planned_end), "teamLead": p.team_lead,
               "weight": p.weight, "scopeNote": p.scope_note}
        ch = [k for k in cur if k in rec and rec[k] != cur[k]]
        if ch:
            out["phaseChanges"].append({"phaseId": p.id, "name": p.name, "fields": ch})
    return out


def batch_dict(b):
    preview = b.result.get("preview") if isinstance(b.result, dict) else None
    return {"id": b.id, "kind": b.kind, "status": b.status, "fileName": b.file.original_name if b.file_id else "",
            "createdAt": iso(b.created_at), "createdBy": b.created_by.name, "errors": b.errors,
            "sheets": b.parsed.get("sheets") if isinstance(b.parsed, dict) else None,
            "taskRows": b.parsed.get("taskRows") if isinstance(b.parsed, dict) else None,
            "preview": preview, "result": b.result.get("done") if isinstance(b.result, dict) else None}


def preview_summary(c):
    return {"add": len(c["add"]), "update": len(c["update"]), "signed": c["signed"], "same": len(c["same"]),
            "phaseChanges": c["phaseChanges"]}


class PlanImports(ProjectView):
    def get(self, request, project_id):
        self.require("plan.import", "只有项目经理可以导入导出计划")
        qs = ImportBatch.objects.filter(project=self.project, kind__in=["plan_xlsx", "plan_csv"]).select_related(
            "file", "created_by").order_by("-id")[:10]
        return Response({"items": [batch_dict(b) for b in qs]})

    def post(self, request, project_id):
        """上传并逐行校验；返回错误清单（工作表 + 行号）或预计变更（AC-F03-08）。"""
        self.require("plan.import", "只有项目经理可以导入导出计划")
        f = request.FILES.get("file")
        if f is None:
            fail("VALIDATION_FAILED", "请选择文件")
        ext = Path(f.name).suffix.lower()
        if ext not in (".xlsx", ".xlsm", ".csv"):
            raise ApiError(415, "FILE_TYPE", "只支持 .xlsx 或 .csv 计划文件，请重新选择")
        if f.size > 20 * 1024 * 1024:
            raise ApiError(413, "FILE_TOO_LARGE", f"{f.name}：超过 20MB，未上传")
        data = f.read()
        try:
            sheets = read_workbook(f.name, data)
        except Exception:  # noqa: BLE001 — 文件损坏或不是表格：明确告知，不入库
            raise ApiError(400, "VALIDATION_FAILED", "文件无法解析，请确认是 .xlsx 或 .csv 计划文件")
        v = Validator(self.project)
        parsed = v.run(sheets)
        sf = storage.save_bytes(data, f.name, request.user, self.project)
        preview = preview_summary(classify_tasks(self.project, parsed)) if not v.errors else None
        b = ImportBatch.objects.create(project=self.project, kind="plan_csv" if ext == ".csv" else "plan_xlsx",
                                       file=sf, parsed=parsed, errors=v.errors, result={"preview": preview},
                                       status="校验失败" if v.errors else "待导入", created_by=request.user)
        audit.log(request.user, "plan.import.validate", "import_batch", b.id, f.name, None,
                  {"errors": len(v.errors)}, project=self.project)
        return Response(batch_dict(ImportBatch.objects.select_related("file", "created_by").get(pk=b.pk)), status=201)


class PlanImportExecute(ProjectView):
    def post(self, request, project_id, pk):
        self.require("plan.import", "只有项目经理可以导入导出计划")
        with transaction.atomic():
            b = ImportBatch.objects.select_for_update().get(pk=self.obj(ImportBatch, pk).pk)
            if b.kind not in ("plan_xlsx", "plan_csv"):
                fail("NOT_FOUND", "记录不存在", 404)
            if b.errors:
                fail("IMPORT_HAS_ERRORS", "文件有错误行，全部修正前不能执行导入", 409)
            if b.status == "已导入":
                fail("ALREADY_IMPORTED", "该文件已导入，请重新上传", 409)
            done = execute_plan(self.project, b, request.user)
            b.status = "已导入"
            b.result = {**b.result, "done": done}
            b.save(update_fields=["status", "result"])
            audit.log(request.user, "plan.import.execute", "import_batch", b.id, b.file.original_name if b.file_id else "",
                      None, {k: v for k, v in done.items() if k != "signed"}, project=self.project)
        return Response(batch_dict(ImportBatch.objects.select_related("file", "created_by").get(pk=b.pk)))


def execute_plan(project, b, actor):
    """单事务：新增工作项、更新未签认工作项、已签认的跳过；有变化的工段计划写留档。调用方需在事务内。"""
    parsed = b.parsed
    c = classify_tasks(project, parsed)
    by_row = {t["row"]: t for t in parsed["tasks"]}
    phases = {p.id: p for p in Phase.objects.filter(project=project)}
    floors = {f.id: f for f in Floor.objects.filter(project=project)}
    rooms = {r.id: r for r in Room.objects.filter(project=project)}
    fname = b.file.original_name if b.file_id else "计划文件"
    for e in c["add"]:
        t = by_row[e["row"]]
        vals = {"phase": phases[t["phaseId"]], "name": t["name"], "floor": floors.get(t.get("floorId")),
                "room": rooms.get(t.get("roomId")), "description": t.get("description", ""),
                "acceptanceCriteria": t.get("acceptanceCriteria", ""),
                "plannedFinish": st.parse_date(t["plannedFinish"]) if t.get("plannedFinish") else None,
                "actualFinish": st.parse_date(t["actualFinish"]) if t.get("actualFinish") else None,
                "ownerName": t.get("ownerName") or "待指派", "weight": t.get("weight", 2),
                "progress": t.get("progress", 100 if t.get("actualFinish") else 0)}
        st.create_work_item(project, vals, actor, source="excel_import")
    skipped = set()
    for e in c["update"]:
        t = by_row[e["row"]]
        w = WorkItem.objects.select_for_update().get(pk=e["workItemId"])
        bl = Baseline.objects.select_for_update().filter(work_item=w).first()
        if bl and bl.status in LOCKED_BASELINE:  # 校验后到执行前被提交或签认：同样不改
            c["signed"].append({**e, "baselineStatus": bl.status, "note": SIGNED_NOTE})
            skipped.add(w.id)
            continue
        if "floorId" in t:
            w.floor = floors.get(t["floorId"])
            w.room = rooms.get(t["roomId"])
        for key, attr in (("plannedFinish", "planned_finish"), ("actualFinish", "actual_finish")):
            if key in t:
                setattr(w, attr, st.parse_date(t[key]))
        for key, attr in (("ownerName", "owner_name"), ("weight", "weight"), ("progress", "progress"),
                          ("description", "description"), ("acceptanceCriteria", "acceptance_criteria")):
            if key in t:
                setattr(w, attr, t[key])
        w.save()
    c["update"] = [e for e in c["update"] if e["workItemId"] not in skipped]
    logs = 0
    for rec in parsed["phases"]:
        p = Phase.objects.select_for_update().get(pk=rec["phaseId"])
        new = {}
        for key in ("plannedStart", "plannedEnd"):
            if key in rec:
                new[key] = st.parse_date(rec[key])
        for key in ("teamLead", "weight", "scopeNote"):
            if key in rec:
                new[key] = rec[key]
        ps, pe = new.get("plannedStart", p.planned_start), new.get("plannedEnd", p.planned_end)
        if ps > pe:
            fail("IMPORT_HAS_ERRORS", f"施工阶段「{p.name}」开始日期不能晚于结束日期", 409)
        if st.apply_phase_plan(p, new, actor, f"Excel 导入：{fname}", source="excel_import"):
            logs += 1
    sup = 0
    titles = set(Supplement.objects.filter(project=project).values_list("title", flat=True))
    today = st.today()
    for x in parsed["extras"]:
        if x["title"] in titles:
            continue
        raised = st.parse_date(x.get("raisedOn")) or today
        Supplement.objects.create(project=project, title=x["title"], note=x["note"], raised_on=raised,
                                  due_on=st.parse_date(x.get("dueOn")) or raised + datetime.timedelta(days=7),
                                  owner_name=x["ownerName"], status=x["status"], source="excel_import")
        titles.add(x["title"])
        sup += 1
    return {"added": len(c["add"]), "updated": len(c["update"]), "signed": c["signed"], "same": len(c["same"]),
            "errorRows": 0, "phaseLogs": logs, "supplementsAdded": sup}


# ---------------------------------------------------------------- 原型 v1.20 JSON 备份导入（F18）

class InvalidBackup(Exception):
    pass


BOARD = {"owner": "owner", "manager": "site", "designer": "design"}
ZONE = {"ld": "layout", "r3": "render", "ct": "contract"}
PROTO_LOG_FIELDS = {"start": ("plannedStart", "计划开始日期"), "end": ("plannedEnd", "计划结束日期"),
                    "crew": ("teamLead", "班组负责人"), "w": ("weight", "工段权重"), "scope": ("scopeNote", "施工范围")}


def _need(cond):
    if not cond:
        raise InvalidBackup()


def check_backup(o, phase_codes):
    """结构校验：必须是原型 exportJSON 导出的 OVERLAY 对象。"""
    _need(isinstance(o, dict))
    _need(all(k in o for k in ("phase", "tasks", "att", "issues")))
    _need(isinstance(o["phase"], dict) and isinstance(o["tasks"], dict) and isinstance(o["att"], dict)
          and isinstance(o["issues"], list))
    for k, typ in (("log", list), ("files", list), ("extra", list), ("notice", list), ("photos", dict),
                   ("closed", dict), ("edit", dict), ("folders", dict)):
        _need(k not in o or o[k] is None or isinstance(o[k], typ))
    for pid in list(o["phase"]) + list(o["tasks"]) + list(o["att"]) + list((o.get("photos") or {})):
        _need(pid in phase_codes)
    for pid, tasks in o["tasks"].items():
        _need(isinstance(tasks, list))
        for t in tasks:
            _need(isinstance(t, dict) and isinstance(t.get("n"), str) and t["n"].strip())
    for pid, days in o["att"].items():
        _need(isinstance(days, dict))
        for k, v in days.items():
            _need(st.parse_date(k) is not None)
            _need(isinstance(v, (int, float)) and 0 <= v <= 40 and float(v).is_integer())
    for i in o["issues"]:
        _need(isinstance(i, dict) and isinstance(i.get("t"), str) and i.get("pid") in phase_codes)


def _num(v, default):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return default


def _missing_file(project, meta, actor):
    """原型资料只有元数据：建 status=missing 的文件记录（「文件缺失，需重新上传」）。"""
    name = str(meta.get("name") or f"资料-{meta.get('id', '')}")[:200]
    return StoredFile.objects.create(project=project, storage_key="", sha256="", size=_num(meta.get("size"), 0),
                                     mime=str(meta.get("type") or "application/octet-stream")[:100],
                                     original_name=name, uploader=actor, status="missing")


def import_prototype(project, o, actor):
    """单事务导入；任何异常由调用方回滚。返回各类计数与缺失文件清单。"""
    phases = {p.code: p for p in Phase.objects.filter(project=project)}
    floors = {f.code: f for f in Floor.objects.filter(project=project)}
    rooms = {r.code: r for r in Room.objects.filter(project=project)}
    files_meta = {str(f.get("id")): f for f in (o.get("files") or []) if isinstance(f, dict)}
    missing = []
    counts = {"workItems": 0, "issues": 0, "attendance": 0, "planLogs": 0, "supplements": 0, "notices": 0,
              "photos": 0}

    # 工段计划字段覆盖（原型的 w 为小数比例）
    for code, ov in o["phase"].items():
        p = phases[code]
        if not isinstance(ov, dict):
            raise InvalidBackup()
        if ov.get("start") and st.parse_date(ov["start"]):
            p.planned_start = st.parse_date(ov["start"])
        if ov.get("end") and st.parse_date(ov["end"]):
            p.planned_end = st.parse_date(ov["end"])
        if "crew" in ov:
            p.team_lead = str(ov["crew"] or "")[:60]
        if "scope" in ov:
            p.scope_note = str(ov["scope"] or "")
        if "w" in ov:
            w = float(ov["w"])
            p.weight = max(1, int(round(w * 100 if w <= 1 else w)))
        if p.planned_start > p.planned_end:
            raise InvalidBackup()
        p.save()

    # 工作事项：存在即整段替换该工段（原型「快照式落库」语义）
    rec_counts = st.item_record_counts(project)
    for code, tasks in o["tasks"].items():
        p = phases[code]
        current = {}
        for w in WorkItem.objects.filter(phase=p, is_disabled=False).order_by("id"):
            current.setdefault(w.name, w)
        keep = set()
        for t in tasks:
            fl = floors.get(t.get("fl"))
            rm = rooms.get(t.get("rm"))
            if rm and (fl is None or rm.floor_id != fl.id):
                rm = None
            vals = {"floor": fl, "room": rm, "description": str(t.get("note") or ""),
                    "acceptance_criteria": str(t.get("std") or ""), "owner_name": str(t.get("owner") or "待指派")[:40],
                    "weight": min(5, max(1, _num(t.get("w"), 2))), "progress": min(100, max(0, _num(t.get("pct"), 0))),
                    "planned_finish": _proto_date(t.get("plan"), p), "actual_finish": _proto_date(t.get("act"), p),
                    "source": "prototype_import"}
            w = current.get(t["n"].strip())
            if w and w.id not in keep:
                for k, v in vals.items():
                    if k in ("description", "acceptance_criteria") and not v:
                        continue
                    setattr(w, k, v)
                w.save()
            else:
                w = WorkItem(project=project, phase=p, name=t["n"].strip()[:120], **vals)
                st.apply_rule(w)
                w.save()
                Baseline.objects.create(project=project, work_item=w)
            keep.add(w.id)
            counts["workItems"] += 1
            for fid in t.get("files") or []:
                meta = files_meta.get(str(fid), {"id": fid, "name": f"资料-{fid}"})
                sf = _missing_file(project, meta, actor)
                WorkItemAttachment.objects.create(project=project, work_item=w, file=sf)
                missing.append({"name": sf.original_name, "location": f"工作项：{w.name}", "workItemId": w.id})
        for w in WorkItem.objects.filter(phase=p, is_disabled=False).exclude(id__in=keep):
            if rec_counts.get(w.id):
                w.is_disabled = True
                w.save(update_fields=["is_disabled"])
            else:
                Baseline.objects.filter(work_item=w).delete()
                w.delete()

    # 人工出勤登记
    for code, days in o["att"].items():
        for k, v in days.items():
            Attendance.objects.update_or_create(project=project, phase=phases[code], date=st.parse_date(k),
                                                defaults={"headcount": int(v), "source": "prototype_import"})
            counts["attendance"] += 1

    # 问题（新增问题 + 字段修正 + 闭环覆盖）
    edit, closed = o.get("edit") or {}, o.get("closed") or {}
    for raw in o["issues"]:
        i = {**raw, **(edit.get(raw.get("id")) or {})}
        p = phases[i["pid"]]
        fl, rm = floors.get(i.get("fl")), rooms.get(i.get("rm"))
        if rm and (fl is None or rm.floor_id != fl.id):
            rm = None
        found = st.parse_date(i.get("d")) or st.today()
        status = "已闭环" if raw.get("id") in closed else (i.get("st") if i.get("st") in ("未闭环", "处理中", "已闭环") else "未闭环")
        issue = Issue.objects.create(
            project=project, seq_no=next_issue_seq(project), phase=p, floor=fl, room=rm,
            description=str(i["t"]), level=i.get("lv") if i.get("lv") in Issue.LEVELS else "一般",
            countermeasure=str(i.get("fix") or "待补充对策"), owner_name=str(i.get("owner") or "待指派")[:40],
            found_on=found, due_date=st.parse_date(i.get("due")), status=status,
            closed_on=st.parse_date(closed.get(raw.get("id"))) if status == "已闭环" else None,
            source="prototype_import", created_by=actor)
        issue_event(issue, actor, "created", after=status, reason="原型导入")
        counts["issues"] += 1

    # 计划变更留档：原记录时间写入 changes.originalAt（关键记录时间不接受客户端传值，NFR-008）
    for r in o.get("log") or []:
        if not isinstance(r, dict) or r.get("pid") not in phases:
            raise InvalidBackup()
        items, delta, has_date = [], 0, False
        for it in r.get("items") or []:
            key, label = PROTO_LOG_FIELDS.get(it.get("f"), (str(it.get("f")), str(it.get("f"))))
            items.append({"field": key, "label": label, "before": str(it.get("from", "")), "after": str(it.get("to", "")),
                          "originalAt": str(r.get("t") or "")})
            if it.get("f") in ("start", "end") and it.get("dd") is not None:
                has_date = True
                delta += _num(it["dd"], 0) * (1 if it["f"] == "end" else -1)
        if not items:
            continue
        log = PlanChangeLog.objects.create(project=project, phase=phases[r["pid"]], actor=actor, changes=items,
                                           duration_delta_days=delta if has_date else None,
                                           reason=str(r.get("note") or "")[:200], source="prototype_import")
        ledger.append(project.id, log)
        counts["planLogs"] += 1

    # 补充事项
    for x in o.get("extra") or []:
        if not isinstance(x, dict) or not str(x.get("title") or "").strip():
            raise InvalidBackup()
        raised = st.parse_date(x.get("date")) or st.today()
        s = Supplement.objects.create(project=project, title=str(x["title"])[:100], note=str(x.get("desc") or ""),
                                      raised_on=raised, due_on=st.parse_date(x.get("due")) or raised + datetime.timedelta(days=7),
                                      owner_name=str(x.get("owner") or "待指派")[:40],
                                      status=x.get("st") if x.get("st") in SUPPLEMENT_STATUS else "待处理",
                                      source="prototype_import")
        for fid in x.get("files") or []:
            sf = _missing_file(project, files_meta.get(str(fid), {"id": fid}), actor)
            s.files.add(sf)
            missing.append({"name": sf.original_name, "location": f"补充事项：{s.title}"})
        counts["supplements"] += 1

    # 公告留言（先根后回复）
    notes = [n for n in (o.get("notice") or []) if isinstance(n, dict)]
    made = {}
    for n in sorted(notes, key=lambda n: 1 if n.get("re") else 0):
        parent = made.get(n.get("re")) if n.get("re") else None
        made[n.get("id")] = Notice.objects.create(project=project, board=BOARD.get(n.get("board"), "owner"),
                                                  author=None, author_name=str(n.get("who") or "原型用户")[:40],
                                                  parent=parent, body=str(n.get("text") or ""),
                                                  source="prototype_import")
        counts["notices"] += 1

    # 验收图片：Base64 内容解码后写入存储
    for code, photos in (o.get("photos") or {}).items():
        if not isinstance(photos, list):
            raise InvalidBackup()
        for ph in photos:
            cap = str(ph.get("cap") or "验收图片")[:100]
            sf = None
            src = str(ph.get("src") or "")
            m = re.match(r"data:(image/[\w.+-]+);base64,(.+)$", src, re.S)
            if m:
                ext = {"image/png": ".png", "image/webp": ".webp"}.get(m[1], ".jpg")
                sf = storage.save_bytes(base64.b64decode(m[2]), cap + ext, actor, project, mime=m[1])
            AcceptancePhoto.objects.create(project=project, phase=phases[code], file=sf, caption=cap,
                                           status=ph.get("st") if ph.get("st") in ("施工中", "待复验", "已验收") else "待复验",
                                           source="prototype_import")
            counts["photos"] += 1

    # 房屋档案资料（只有元数据）
    folders = o.get("folders") or {}
    for f in o.get("files") or []:
        if not isinstance(f, dict) or f.get("sec") not in ZONE:
            continue
        zone = ZONE[f["sec"]]
        fname = next((x.get("name") for x in folders.get(f["sec"], []) if isinstance(x, dict) and x.get("id") == f.get("fo")),
                     None) or "原型导入"
        folder = ArchiveFolder.objects.filter(project=project, zone=zone, name=fname).first() or \
            ArchiveFolder.objects.create(project=project, zone=zone, name=str(fname)[:40], order_no=99)
        sf = _missing_file(project, f, actor)
        ArchiveFile.objects.create(project=project, folder=folder, file=sf)
        missing.append({"name": sf.original_name, "location": f"房屋档案 · {folder.name}", "zone": zone})
    return {"counts": counts, "missing": missing}


def _proto_date(v, phase):
    if not v:
        return None
    s = str(v)
    if re.fullmatch(r"\d{1,2}-\d{1,2}", s):
        try:
            return _date_in_year(s, phase.planned_start, phase.planned_end)
        except ValueError:
            return None
    return st.parse_date(s)


def proto_dict(b):
    r = b.result if isinstance(b.result, dict) else {}
    return {"id": b.id, "status": b.status, "fileName": b.file.original_name if b.file_id else "",
            "createdAt": iso(b.created_at), "createdBy": b.created_by.name, "projectStatus": r.get("projectStatus"),
            "counts": r.get("counts"), "missing": r.get("missing", []), "message": r.get("message")}


class PrototypeImports(ProjectView):
    def get(self, request, project_id):
        if not (self.can("prototype.import") or self.can("plan.import")):
            fail("FORBIDDEN", "无权查看原型备份导入记录", 403)
        qs = ImportBatch.objects.filter(project=self.project, kind="prototype_json").select_related(
            "file", "created_by").order_by("-id")
        return Response({"items": [proto_dict(b) for b in qs], "projectStatus": self.project.status})

    def post(self, request, project_id):
        self.require("prototype.import", "只有公司管理员可以导入原型备份")
        if self.project.status != "筹备":
            fail("PROJECT_NOT_PREPARING", f"当前项目状态为「{self.project.status}」。原型备份只能导入「筹备」状态的项目", 409)
        f = request.FILES.get("file")
        if f is None:
            fail("VALIDATION_FAILED", "请选择文件")
        data = f.read()
        sf = storage.save_bytes(data, f.name, request.user, self.project)
        phase_codes = set(Phase.objects.filter(project=self.project).values_list("code", flat=True))
        status_then = self.project.status
        try:
            o = json.loads(data.decode("utf-8-sig"))
            check_backup(o, phase_codes)
            with transaction.atomic():
                out = import_prototype(self.project, o, request.user)
        except Exception:  # noqa: BLE001 — 任何解析或结构错误都整单回滚并拒绝（AC-F18-03）
            b = ImportBatch.objects.create(project=self.project, kind="prototype_json", file=sf, status="已拒绝",
                                           errors=[{"message": "不是有效的原型备份文件"}],
                                           result={"message": "不是有效的原型备份文件。项目数据未改变。",
                                                   "projectStatus": status_then}, created_by=request.user)
            audit.log(request.user, "prototype.import.reject", "import_batch", b.id, f.name, project=self.project)
            raise ApiError(422, "INVALID_PROTOTYPE_BACKUP", "不是有效的原型备份文件", details={"batchId": b.id})
        b = ImportBatch.objects.create(project=self.project, kind="prototype_json", file=sf, status="已导入",
                                       result={**out, "projectStatus": status_then}, created_by=request.user)
        audit.log(request.user, "prototype.import", "import_batch", b.id, f.name, None, out["counts"],
                  project=self.project)
        return Response(proto_dict(ImportBatch.objects.select_related("file", "created_by").get(pk=b.pk)), status=201)


class PrototypeImportDetail(ProjectView):
    def get(self, request, project_id, pk):
        b = self.obj(ImportBatch, pk, kind="prototype_json")
        return Response(proto_dict(b))


P = "projects/<int:project_id>/"
urlpatterns = [
    path(P + "plan/template.xlsx", PlanTemplate.as_view()),
    path(P + "plan/export.xlsx", PlanExport.as_view()),
    path(P + "plan/imports", PlanImports.as_view()),
    path(P + "plan/imports/<int:pk>/execute", PlanImportExecute.as_view()),
    path(P + "prototype-imports", PrototypeImports.as_view()),
    path(P + "prototype-imports/<int:pk>", PrototypeImportDetail.as_view()),
]
