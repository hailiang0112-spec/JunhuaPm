"""archive 代理验收：FLOW-17（档案部分）、FLOW-18、FLOW-19（CSV 部分）、FLOW-20、FLOW-22、RULE-02。

前置数据（核对、放行、签认等关键记录）直接用模型构造并 ledger.append 入链；期望值取自契约中的已知业务数据。
"""
import csv
import datetime
import hashlib
import importlib
import io
import zipfile
from unittest import mock

import psycopg
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image

from core.models import (AuditLog, Baseline, BaselineRevision, BaselineSignoff, ChangeOrder, Checkpoint, CheckPhoto, CheckRecord,
                         DrawingSet, DrawingVersion, Issue, LedgerEntry, Phase, PlanChangeLog, ReleaseRecord, StoredFile,
                         Supplement, WorkItem)
from core.services import ledger, storage

TZ = timezone.get_default_timezone()


def at(y, m, d, hh=10, mm=0):
    return datetime.datetime(y, m, d, hh, mm, tzinfo=TZ)


def png(name, color=(200, 120, 40)):
    buf = io.BytesIO()
    Image.new("RGB", (64, 48), color).save(buf, "PNG")
    return SimpleUploadedFile(name, buf.getvalue(), content_type="image/png")


def pdf(name, body="现场复尺记录"):
    return SimpleUploadedFile(name, ("%PDF-1.4\n% " + body + "\n%%EOF").encode(), content_type="application/pdf")


def append(project, obj):
    with transaction.atomic():
        obj.save()
        ledger.append(project.id, obj)
    return obj


def frozen(when):
    return mock.patch("django.utils.timezone.now", return_value=when.astimezone(datetime.timezone.utc))


class Evidence:
    """在指定工段造检查点 / 核对 / 照片 / 放行等关键记录。"""

    def __init__(self, project, users):
        self.p, self.u = project, users
        sf = storage.save_bytes(b"%PDF-1.4\n% shuidian v1\n%%EOF", "水电点位图-V1.pdf", users["design_manager"], project)
        self.ds = DrawingSet.objects.create(project=project, name="水电点位图")
        with frozen(at(2026, 9, 20)):
            self.dv = append(project, DrawingVersion(project=project, drawing_set=self.ds, version_no=1, file=sf,
                                                     created_by=users["design_manager"], change_note="首版"))
        self.seq = 200

    def wi(self, phase_name, i=0):
        return WorkItem.objects.filter(project=self.p, phase__name=phase_name).order_by("id")[i]

    def checkpoint(self, wi, status="未核对"):
        self.seq += 1
        return Checkpoint.objects.create(project=self.p, work_item=wi, code=f"CK-{self.seq}", point="线管走向与点位图一致",
                                         drawing_version=self.dv, x=0.4, y=0.5, status=status, created_by=self.u["qc"])

    def check(self, cp, when, conclusion="一致", photo=True, note=""):
        with frozen(when):
            cr = append(self.p, CheckRecord(project=self.p, checkpoint=cp, checker=self.u["qc"], drawing_version=self.dv,
                                            was_latest_signed=True, conclusion=conclusion, was_covered=False, note=note,
                                            discovery_phase="覆盖前" if conclusion == "不一致" else ""))
            ph = None
            if photo:
                sf = storage.save_bytes(f"photo-{cp.code}-{cr.id}".encode() + b"\xff\xd8", f"{cp.code}-现场.jpg", self.u["qc"], self.p)
                ph = append(self.p, CheckPhoto(project=self.p, check_record=cr, file=sf, file_sha256=sf.sha256, pin_x=0.4, pin_y=0.5))
        cp.latest_check = cr
        cp.status = "不一致" if conclusion == "不一致" else "待放行"
        cp.save()
        return cr, ph

    def release(self, cp, cr, when):
        with frozen(when):
            rr = append(self.p, ReleaseRecord(project=self.p, checkpoint=cp, releaser=self.u["owner"], is_proxy=False, basis_check=cr))
        cp.status = "已放行"
        cp.save()
        return rr

    def baseline_sign(self, wi, when):
        bl = Baseline.objects.get(work_item=wi)
        with frozen(when):
            rev = append(self.p, BaselineRevision(project=self.p, baseline=bl, rev_no=1, kind="submit", actor=self.u["pm"],
                                                  content={"name": wi.name, "drawingVersionId": self.dv.id, "cost": "1800.00"}))
            so = append(self.p, BaselineSignoff(project=self.p, revision=rev, signer=self.u["owner"],
                                                content_snapshot={"工作项": wi.name, "费用": "1800.00", "图纸": "水电点位图 V1"}))
        bl.status, bl.current_revision, bl.signed_revision = "已签认", rev, rev
        bl.save()
        return so


# ———————————————————————————— FLOW-17 档案部分 ————————————————————————————

def folders(c, pid, zone):
    return {f["name"]: f for f in c.get(f"/api/projects/{pid}/archive/{zone}/folders").json()["items"]}


def test_flow17_archive_folders_upload_void_and_contract(project, client_for, users):
    pid = project.id
    pm, qc = client_for("pm"), client_for("qc")
    layout = folders(pm, pid, "layout")
    assert len(layout) == 6 and len(folders(pm, pid, "contract")) == 4

    # 分类管理：空名称报错、新建、重名、改名；非项目经理 / 公司管理员不能管理
    r = pm.post(f"/api/projects/{pid}/archive/layout/folders", {"name": " "}, format="json")
    assert r.status_code == 400 and r.json()["fields"] == [{"field": "name", "message": "请填写分类名称"}]
    r = pm.post(f"/api/projects/{pid}/archive/layout/folders", {"name": "地下室车库"}, format="json")
    assert r.status_code == 201 and r.json()["fileCount"] == 0
    garage = r.json()["id"]
    r = pm.post(f"/api/projects/{pid}/archive/layout/folders", {"name": "地下室车库"}, format="json")
    assert r.status_code == 400 and r.json()["fields"][0]["message"] == "已有同名分类"
    r = pm.patch(f"/api/projects/{pid}/archive/layout/folders/{garage}", {"name": "负一楼车库"}, format="json")
    assert r.status_code == 200 and r.json()["name"] == "负一楼车库"
    assert client_for("qc").post(f"/api/projects/{pid}/archive/layout/folders", {"name": "随便"}, format="json").status_code == 403
    assert client_for("admin").post(f"/api/projects/{pid}/archive/layout/folders", {"name": "竣工图"}, format="json").status_code == 201

    # 上传：2 张图片 + 1 个 PDF 合格，60MB 文件单独被拒并给出文件名与原因
    target = list(layout.values())[2]  # 第三个默认分类（一楼）
    url = f"/api/projects/{pid}/archive/layout/folders/{target['id']}/files"
    ok = [qc.post(url, {"file": f}, format="multipart") for f in (png("一楼客厅墙面.png"), png("一楼厨房管线.png", (10, 90, 160)),
                                                                     pdf("一楼现场复尺记录.pdf"))]
    assert [r.status_code for r in ok] == [201, 201, 201]
    big = SimpleUploadedFile("一楼全景视频.pdf", b"0" * (60 * 1024 * 1024), content_type="application/pdf")
    r = qc.post(url, {"file": big}, format="multipart")
    assert r.status_code == 413 and r.json()["message"] == "一楼全景视频.pdf：超过 50MB，未上传"
    listing = qc.get(url).json()
    assert listing["folder"]["fileCount"] == 3 and len(listing["items"]) == 3
    imgs = [f for f in listing["items"] if f["isImage"]]
    assert len(imgs) == 2 and all(f["thumbUrl"] for f in imgs)
    pdf_item = next(f for f in listing["items"] if f["name"] == "一楼现场复尺记录.pdf")
    dl = qc.get(pdf_item["url"])
    assert dl.status_code == 200 and dl.content.startswith(b"%PDF")
    assert pdf_item["uploaderName"] == "王工" and pdf_item["canVoid"] is True

    # 作废：普通资料作废后默认隐藏并留痕；他人不能作废；被核对记录引用的照片作废被拒
    assert pm.post(f"/api/projects/{pid}/archive-files/{pdf_item['archiveFileId']}/void").status_code == 403
    r = qc.post(f"/api/projects/{pid}/archive-files/{pdf_item['archiveFileId']}/void")
    assert r.status_code == 200 and r.json()["status"] == "voided" and r.json()["voidedByName"] == "王工"
    assert len(qc.get(url).json()["items"]) == 2
    shown = qc.get(url + "?showVoided=1").json()["items"]
    assert len(shown) == 3 and next(f for f in shown if f["id"] == pdf_item["id"])["voidedAt"]
    log = AuditLog.objects.get(project=project, action="作废", object_id=str(pdf_item["id"]))
    assert log.actor == users["qc"] and log.after == {"status": "已作废"}
    assert qc.post(f"/api/projects/{pid}/archive-files/{pdf_item['archiveFileId']}/void").status_code == 409

    ev = Evidence(project, users)
    cp = ev.checkpoint(ev.wi("水电"))
    cr, _ = ev.check(cp, timezone.now(), photo=False)
    photo = StoredFile.objects.get(pk=imgs[0]["id"])
    append(project, CheckPhoto(project=project, check_record=cr, file=photo, file_sha256=photo.sha256, pin_x=0.1, pin_y=0.2))
    r = qc.post(f"/api/projects/{pid}/archive-files/{imgs[0]['archiveFileId']}/void")
    assert r.status_code == 409 and r.json()["message"] == "该文件已被核对、签认或变更单引用，不能作废"
    assert StoredFile.objects.get(pk=photo.id).status == "normal"

    # 删除分类：非空分类提示先移出文件；空分类可删；移动后原分类变空
    r = pm.delete(f"/api/projects/{pid}/archive/layout/folders/{target['id']}")
    assert r.status_code == 409 and r.json()["message"] == f"分类「{target['name']}」下还有 3 个文件，请先移出文件"
    empty = folders(pm, pid, "layout")["夹层"]
    assert pm.delete(f"/api/projects/{pid}/archive/layout/folders/{empty['id']}").status_code == 204
    assert "夹层" not in folders(pm, pid, "layout")
    r = pm.post(f"/api/projects/{pid}/archive-files/{imgs[1]['archiveFileId']}/move", {"folderId": garage}, format="json")
    assert r.status_code == 200 and r.json()["folderId"] == garage

    # 合同专区：设计经理、质检被拒；业主、项目经理、预算专员、公司管理员可见
    contract = folders(pm, pid, "contract")["主合同"]
    r = client_for("budget").post(f"/api/projects/{pid}/archive/contract/folders/{contract['id']}/files",
                                  {"file": pdf("君华新城联排别墅装修工程施工合同.pdf", "合同")}, format="multipart")
    assert r.status_code == 201
    contract_file = r.json()
    for key in ("design_manager", "qc"):
        c = client_for(key)
        for u in (f"/api/projects/{pid}/archive/contract/folders", f"/api/projects/{pid}/archive/contract/folders/{contract['id']}/files",
                  contract_file["url"]):
            r = c.get(u)
            assert r.status_code == 403 and r.json()["message"] == "无权查看合同资料", (key, u)
    for key in ("owner", "pm", "budget", "admin"):
        c = client_for(key)
        assert c.get(f"/api/projects/{pid}/archive/contract/folders").status_code == 200
        items = c.get(f"/api/projects/{pid}/archive/contract/folders/{contract['id']}/files").json()["items"]
        assert [f["name"] for f in items] == ["君华新城联排别墅装修工程施工合同.pdf"]
        assert c.get(contract_file["url"]).status_code == 200


# ———————————————————————————— FLOW-18 补充事项与公告栏 ————————————————————————————

def test_flow18_supplements(project, client_for, users):
    pid = project.id
    pm = client_for("pm")
    r = pm.post(f"/api/projects/{pid}/supplements", {"title": "", "note": "可空"}, format="json")
    assert r.status_code == 400
    assert {f["field"]: f["message"] for f in r.json()["fields"]} == {
        "title": "请填写标题", "raisedOn": "请选择提出日期", "dueOn": "请选择约定完成日期", "ownerName": "请选择责任人"}
    r = pm.post(f"/api/projects/{pid}/supplements", {"title": "二楼次卧加装窗帘盒", "raisedOn": "2026-10-26", "dueOn": "2026-10-20",
                                                     "ownerName": "张工"}, format="json")
    assert r.status_code == 400 and r.json()["fields"] == [{"field": "dueOn", "message": "约定完成日期不能早于提出日期"}]

    att = pm.post(f"/api/projects/{pid}/files", {"file": png("浴缸位置示意.png")}, format="multipart").json()
    r = pm.post(f"/api/projects/{pid}/supplements", {
        "title": "三楼卫生间增加浴缸", "note": "业主现场提出在三楼主卫增加独立浴缸，需预留给排水。", "raisedOn": "2026-10-27",
        "dueOn": "2026-11-05", "ownerName": "张工", "fileIds": [att["id"]]}, format="json")
    assert r.status_code == 201
    tub = r.json()
    assert tub["status"] == "待处理" and tub["code"] == "BC-001" and [f["name"] for f in tub["files"]] == ["浴缸位置示意.png"]
    other = pm.post(f"/api/projects/{pid}/supplements", {"title": "业主自购卫浴到货时间确认", "raisedOn": "2026-10-05",
                                                         "dueOn": "2026-10-15", "ownerName": "刘先生"}, format="json").json()
    late = pm.post(f"/api/projects/{pid}/supplements", {"title": "三楼露台加装户外洗手池", "raisedOn": "2026-10-25",
                                                        "dueOn": "2026-11-10", "ownerName": "张工"}, format="json").json()

    # 状态：待处理 → 处理中 → 已完成，任一成员可切换；已完成排在所有未完成之后
    qc = client_for("qc")
    for st in ("处理中", "已完成"):
        r = qc.post(f"/api/projects/{pid}/supplements/{other['id']}/status", {"status": st}, format="json")
        assert r.status_code == 200 and r.json()["status"] == st
    assert qc.post(f"/api/projects/{pid}/supplements/{other['id']}/status", {"status": "作废"}, format="json").status_code == 400
    r = pm.post(f"/api/projects/{pid}/supplements/{tub['id']}/status", {"status": "处理中"}, format="json")
    assert r.json()["status"] == "处理中"
    listing = pm.get(f"/api/projects/{pid}/supplements").json()
    assert [x["code"] for x in listing["items"]] == ["BC-001", "BC-003", "BC-002"]  # 未完成按约定完成日期升序，已完成在后
    assert listing["counts"] == {"undone": 2, "done": 1}
    assert AuditLog.objects.filter(project=project, object_type="补充事项", action="状态变更").count() == 3

    # 转为变更单：仅项目经理；生成预填内容与附件的草稿，并在补充事项上关联编号
    r = qc.post(f"/api/projects/{pid}/supplements/{tub['id']}/to-change")
    assert r.status_code == 403
    r = pm.post(f"/api/projects/{pid}/supplements/{tub['id']}/to-change")
    assert r.status_code == 201
    linked = r.json()["changeOrder"]
    co = ChangeOrder.objects.get(pk=linked["id"])
    assert linked["code"] == co.code == "BG-001" and co.status == "草稿" and co.title == "三楼卫生间增加浴缸"
    assert "业主现场提出在三楼主卫增加独立浴缸" in co.content and "补充事项 BC-001" in co.content
    assert list(co.files.values_list("id", flat=True)) == [att["id"]] and co.source_supplement_id == tub["id"]
    item = next(x for x in pm.get(f"/api/projects/{pid}/supplements").json()["items"] if x["id"] == tub["id"])
    assert item["changeOrder"]["code"] == "BG-001" and item["canConvert"] is False
    assert pm.post(f"/api/projects/{pid}/supplements/{tub['id']}/to-change").status_code == 409
    assert pm.post(f"/api/projects/{pid}/supplements/{other['id']}/to-change").status_code == 409  # 已完成不能转（DEC-132）
    assert Supplement.objects.get(pk=late["id"]).change_order is None


def test_flow18_notices(project, client_for, users):
    pid = project.id
    owner, pm, qc, dm = client_for("owner"), client_for("pm"), client_for("qc"), client_for("design_manager")
    assert owner.post(f"/api/projects/{pid}/notices", {"body": "  "}, format="json").json()["fields"][0]["message"] == "请填写留言内容"
    with frozen(at(2026, 10, 26, 20, 15)):
        o1 = owner.post(f"/api/projects/{pid}/notices", {"body": "三楼露台想加一个户外洗手池，麻烦张工评估一下给排水能不能预留。"},
                        format="json").json()
    with frozen(at(2026, 10, 27, 7, 50)):
        p1 = pm.post(f"/api/projects/{pid}/notices", {"body": "今天砌墙班组 6 人进场，重点做二楼书房 B 和主卫门洞。"}, format="json").json()
    with frozen(at(2026, 10, 25, 17, 30)):
        q1 = qc.post(f"/api/projects/{pid}/notices", {"body": "砌墙阶段已完成 18 个检查点核对，3 处不一致已生成问题。"}, format="json").json()
    with frozen(at(2026, 10, 27, 8, 55)):
        r2 = dm.post(f"/api/projects/{pid}/notices/{o1['id']}/replies", {"body": "位置建议放在东北角，离雨水立管近。"}, format="json")
    with frozen(at(2026, 10, 26, 20, 40)):
        r1 = pm.post(f"/api/projects/{pid}/notices/{o1['id']}/replies", {"body": "可以预留，明天到现场看位置后登记补充事项。"}, format="json")
    assert r1.status_code == r2.status_code == 201
    assert (o1["board"], p1["board"], q1["board"]) == ("owner", "site", "site")
    assert (o1["authorName"], p1["authorName"], q1["authorName"]) == ("刘先生", "张工", "王工")

    data = owner.get(f"/api/projects/{pid}/notices").json()
    assert [n["authorName"] for n in data["boards"]["owner"]] == ["刘先生"]
    assert [n["id"] for n in data["boards"]["site"]] == [p1["id"], q1["id"]]  # 倒序
    assert data["boards"]["design"] == []  # 回复挂在原留言下，不进设计师板块
    assert [(x["authorName"], x["roleLabel"]) for x in data["boards"]["owner"][0]["replies"]] == [("张工", "项目经理"), ("李工", "设计经理")]
    assert data["me"] == {"name": "刘先生", "roleLabel": "业主", "board": "owner", "boardName": "业主"}
    assert dm.post(f"/api/projects/{pid}/notices", {"body": "二楼平面布局图 V4 已发布。"}, format="json").json()["board"] == "design"
    # 兼任：业主 > 设计经理 > 其他（DEC-133）
    from core.models import ProjectMemberRole
    ProjectMemberRole.objects.create(project=project, user=users["pm"], role="design_manager")
    assert pm.post(f"/api/projects/{pid}/notices", {"body": "兼任设计经理后发布"}, format="json").json()["board"] == "design"
    ProjectMemberRole.objects.create(project=project, user=users["pm"], role="owner")
    assert pm.post(f"/api/projects/{pid}/notices", {"body": "兼任业主后发布"}, format="json").json()["board"] == "owner"

    # 删除：只能删自己的，删除后显示「该留言已删除」并保留删除记录
    r = owner.delete(f"/api/projects/{pid}/notices/{p1['id']}")
    assert r.status_code == 403 and r.json()["message"] == "只能删除自己的留言"
    r = owner.delete(f"/api/projects/{pid}/notices/{o1['id']}")
    assert r.status_code == 200
    gone = next(n for n in owner.get(f"/api/projects/{pid}/notices").json()["boards"]["owner"] if n["id"] == o1["id"])
    assert gone["deleted"] is True and gone["body"] == "" and gone["deletedBy"] == "刘先生" and gone["deletedAt"]
    assert len(gone["replies"]) == 2  # 已有回复保留
    assert AuditLog.objects.filter(project=project, object_type="公告栏留言", action="删除", object_id=str(o1["id"]),
                                   actor=users["owner"]).exists()
    assert owner.post(f"/api/projects/{pid}/notices/{o1['id']}/replies", {"body": "再回复"}, format="json").status_code == 409


# ———————————————————————————— FLOW-20 付款节点 ————————————————————————————

def test_flow20_payment_milestone(project, client_for, users):
    pid = project.id
    budget = client_for("budget")
    water = Phase.objects.get(project=project, name="水电")
    ev = Evidence(project, users)
    # 水电工段：5 个未放行检查点、2 个未闭环不一致问题（另有 1 个已放行、1 个已闭环、1 个手工问题与其他工段数据不计入）
    for st in ("未核对", "未核对", "待放行", "待放行", "不一致"):
        ev.checkpoint(ev.wi("水电", 1), st)
    ev.checkpoint(ev.wi("水电", 2), "已放行")
    ev.checkpoint(ev.wi("砌墙"), "未核对")
    bad_cp = ev.checkpoint(ev.wi("水电", 3), "已放行")
    cr, _ = ev.check(bad_cp, timezone.now(), "不一致", photo=False)
    bad_cp.status = "已放行"
    bad_cp.save()
    for i, st in enumerate(("未闭环", "处理中", "已闭环")):
        Issue.objects.create(project=project, seq_no=i + 1, source_check=cr, checkpoint=bad_cp, work_item=bad_cp.work_item, phase=water,
                             description="水电线管与点位图不一致", found_on=datetime.date(2026, 10, 20), status=st)
    Issue.objects.create(project=project, seq_no=9, phase=water, description="手工登记", found_on=datetime.date(2026, 10, 20))

    r = client_for("pm").post(f"/api/projects/{pid}/payments", {"name": "水电完工款"}, format="json")
    assert r.status_code == 403
    r = budget.post(f"/api/projects/{pid}/payments", {"name": ""}, format="json")
    assert {f["field"] for f in r.json()["fields"]} == {"name", "amount", "plannedOn", "phaseIds"}
    r = budget.post(f"/api/projects/{pid}/payments", {"name": "水电完工款", "amount": 253600, "plannedOn": "2027-01-15",
                                                      "phaseIds": [water.id]}, format="json")
    assert r.status_code == 201
    m = r.json()
    assert m["status"] == "未付" and m["amount"] == 253600 and [p["name"] for p in m["phases"]] == ["水电"]
    listed = client_for("owner").get(f"/api/projects/{pid}/payments").json()["items"][0]
    assert (listed["unreleased"], listed["openIssues"]) == (5, 2)

    # 标记已付：先得到确认提示（相当于取消，状态不变），带 confirmUnreleased=5 确认后已付并留痕
    assert client_for("pm").post(f"/api/projects/{pid}/payments/{m['id']}/mark-paid").status_code == 403
    r = budget.post(f"/api/projects/{pid}/payments/{m['id']}/mark-paid", {}, format="json")
    assert r.status_code == 409 and r.json()["message"] == "还有 5 个检查点未放行，确认标记已付？"
    assert r.json()["details"]["unreleased"] == 5
    assert budget.get(f"/api/projects/{pid}/payments/{m['id']}").json()["status"] == "未付"
    r = budget.post(f"/api/projects/{pid}/payments/{m['id']}/mark-paid", {"confirmUnreleased": 5}, format="json")
    assert r.status_code == 200
    paid = r.json()
    assert paid["status"] == "已付" and paid["paidBy"] == "赵会计" and paid["paidConfirmedUnreleased"] == 5 and paid["paidAt"]
    log = AuditLog.objects.get(project=project, object_type="付款节点", action="状态变更")
    assert log.before == {"status": "未付"} and log.after["status"] == "已付" and log.after["确认时未放行检查点"] == 5
    r = budget.post(f"/api/projects/{pid}/payments/{m['id']}/mark-paid", {"confirmUnreleased": 5}, format="json")
    assert r.status_code == 409 and r.json()["details"]["by"] == "赵会计"


# ———————————————————————————— FLOW-19 CSV 部分 ————————————————————————————

def test_flow19_plan_log_csv(project, client_for, users):
    brick = Phase.objects.get(project=project, name="砌墙")
    PlanChangeLog.objects.create(project=project, phase=brick, actor=users["pm"], duration_delta_days=21, reason="砌块到货延迟三周",
                                 changes=[{"field": "planned_end", "label": "结束日期", "before": "2026-10-18", "after": "2026-11-08"}])
    r = client_for("qc").get(f"/api/projects/{project.id}/plan-logs/export.csv")
    assert r.status_code == 200 and r["Content-Type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0] == ["记录时间", "工段", "变更项", "变更前", "变更后", "工期影响（天）", "变更原因备注"]
    assert rows[1][1:] == ["砌墙", "结束日期", "2026-10-18", "2026-11-08", "+21", "砌块到货延迟三周"]
    assert len(rows) == 2
    assert client_for("qc").get(f"/api/projects/{project.id}/audit-summary").json()["planLogCount"] == 1


# ———————————————————————————— FLOW-22 证据包与操作日志 ————————————————————————————

def test_flow22_evidence_export_and_audit_logs(project, client_for, users):
    pid = project.id
    ev = Evidence(project, users)
    water_wi = ev.wi("水电", 1)
    water_wi.drawing_version = ev.dv
    water_wi.save()
    ev.baseline_sign(water_wi, at(2026, 10, 2))                      # 修订 + 签认：2 条
    a = ev.checkpoint(water_wi)
    cra, _ = ev.check(a, at(2026, 10, 5))                            # 核对 + 照片：2 条
    b = ev.checkpoint(ev.wi("水电", 2))
    ev.check(b, at(2026, 10, 8))                                     # 核对 + 照片：2 条
    ev.release(a, cra, at(2026, 10, 9))                              # 放行：1 条
    c = ev.checkpoint(ev.wi("水电", 3))
    ev.check(c, at(2026, 9, 25))                                     # 时间范围外
    d = ev.checkpoint(ev.wi("砌墙"))
    ev.check(d, at(2026, 10, 10))                                    # 工段范围外
    water = Phase.objects.get(project=project, name="水电")

    for key in ("pm", "design_manager", "qc", "budget"):
        r = client_for(key).post(f"/api/projects/{pid}/evidence-exports", {"phaseId": water.id}, format="json")
        assert r.status_code == 403 and r.json()["message"] == "仅业主和公司管理员可导出证据包"
    owner = client_for("owner")
    r = owner.post(f"/api/projects/{pid}/evidence-exports", {"from": "2026-10-27", "to": "2026-10-01"}, format="json")
    assert r.status_code == 400 and r.json()["fields"][0]["message"] == "开始日期不能晚于结束日期"

    r = owner.post(f"/api/projects/{pid}/evidence-exports", {"phaseId": water.id, "from": "2026-10-01", "to": "2026-10-27"}, format="json")
    assert r.status_code == 201
    job = r.json()
    tail = LedgerEntry.objects.filter(project=project).order_by("-seq").first().hash
    assert job["status"] == "完成" and job["result"]["count"] == 7 and job["result"]["tail"] == tail
    assert job["result"]["fileName"].startswith("君华新城联排别墅-证据包-水电-")
    dl = owner.get(f"/api/projects/{pid}/evidence-exports/{job['id']}/download")
    assert dl.status_code == 200 and dl["Content-Type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(b"".join(dl.streaming_content)))
    names = z.namelist()
    wb = load_workbook(io.BytesIO(z.read("记录清单.xlsx")))
    rows = list(wb["记录清单"].iter_rows(values_only=True))
    assert rows[0][:3] == ("编号", "记录类型", "时间") and len(rows) - 1 == 7
    assert sorted(r[1] for r in rows[1:]) == sorted(["基线修订", "基线签认", "核对记录", "核对照片", "核对记录", "核对照片", "放行记录"])
    photos = [n for n in names if n.startswith("照片原件/")]
    assert len(photos) == 2 and all(n.split("/")[1].startswith(("CK-", )) for n in photos)
    for n in photos:
        ph = CheckPhoto.objects.get(file__original_name=n.split("-KR-")[1].split("-", 1)[1])
        assert hashlib.sha256(z.read(n)).hexdigest() == ph.file_sha256  # 照片原件不改动
    assert [n for n in names if n.startswith("图纸版本/")] == ["图纸版本/水电点位图-V1-水电点位图-V1.pdf"]
    snaps = [n for n in names if n.startswith("签认快照/")]
    assert len(snaps) == 1 and "1800.00" in z.read(snaps[0]).decode()
    digest_rows = list(csv.reader(io.StringIO(z.read("摘要清单.csv").decode("utf-8-sig"))))
    assert digest_rows[0] == ["导出时链尾摘要", tail]
    assert sum(1 for r in digest_rows if len(r) > 6 and r[6] == "是") == 7
    assert z.testzip() is None

    # 全项目、不限时间：链上全部记录都在范围内
    r = owner.post(f"/api/projects/{pid}/evidence-exports", {}, format="json").json()
    assert r["result"]["count"] == LedgerEntry.objects.filter(project=project).count()

    # 操作日志：按时间、对象、操作人筛选，显示前后值；仅业主和公司管理员
    pm = client_for("pm")
    s = pm.post(f"/api/projects/{pid}/supplements", {"title": "一楼入户门槛石改大理石", "raisedOn": "2026-10-26", "dueOn": "2026-11-02",
                                                     "ownerName": "张工"}, format="json").json()
    pm.post(f"/api/projects/{pid}/supplements/{s['id']}/status", {"status": "处理中"}, format="json")
    assert pm.get(f"/api/projects/{pid}/audit-logs").status_code == 403
    logs = owner.get(f"/api/projects/{pid}/audit-logs", {"objectType": "补充事项", "objectId": "BC-001"}).json()
    assert logs["total"] == 2
    top = logs["items"][0]
    assert (top["actorName"], top["actorRole"], top["action"], top["before"], top["after"]) == (
        "张工", "项目经理", "状态变更", {"状态": "待处理"}, {"状态": "处理中"})
    assert top["at"].endswith("+08:00")
    today = timezone.localdate().isoformat()
    assert owner.get(f"/api/projects/{pid}/audit-logs", {"from": today, "to": today, "actorId": users["pm"].id}).json()["total"] == 2
    assert owner.get(f"/api/projects/{pid}/audit-logs", {"to": "2026-01-01"}).json()["total"] == 0
    facets = owner.get(f"/api/projects/{pid}/audit-logs").json()["facets"]
    assert "补充事项" in facets["objectTypes"] and {"id": users["owner"].id, "name": "刘先生"} in facets["actors"]


# ———————————————————————————— RULE-02 关键记录只追加 + 摘要链 ————————————————————————————

KEY_TABLES = importlib.import_module("core.migrations.0002_runtime_grants").KEY_TABLES


def test_rule02_append_only_and_integrity(project, client_for, users):
    pid = project.id
    ev = Evidence(project, users)
    wi = ev.wi("水电", 1)
    cp = ev.checkpoint(wi)
    cr, ph = ev.check(cp, at(2026, 10, 5), note="线管走向一致")
    rr = ev.release(cp, cr, at(2026, 10, 6))
    so = ev.baseline_sign(wi, at(2026, 10, 2))
    log = PlanChangeLog.objects.create(project=project, phase=wi.phase, actor=users["pm"], changes=[], duration_delta_days=0)

    # 1) 每条关键记录保存包含上一条摘要的 SHA-256（独立重算）
    prev = "0" * 64
    entries = list(LedgerEntry.objects.filter(project=project).order_by("seq"))
    assert len(entries) == 6  # 图纸版本、核对、照片、放行、基线修订、基线签认
    for e in entries:
        assert e.prev_hash == prev and e.hash == hashlib.sha256(e.content.encode() + prev.encode()).hexdigest()
        prev = e.hash

    # 2) 接口：公司管理员（及业主）对关键记录发 PUT / PATCH / DELETE 一律 405，原记录不变
    paths = [f"checkpoints/{cp.id}/checks/{cr.id}", f"drawing-versions/{ev.dv.id}", f"releases/{rr.id}",
             f"baseline-signoffs/{so.id}", "change-orders/1/events/1", "issues/1/events/1", f"plan-logs/{log.id}",
             f"audit-logs/{AuditLog.objects.create(project=project, action='新增', object_type='测试').id}"]
    for who in ("admin", "owner"):
        c = client_for(who)
        for p in paths:
            for method in ("put", "patch", "delete"):
                r = getattr(c, method)(f"/api/projects/{pid}/{p}", {"note": "改掉"}, format="json")
                assert r.status_code == 405 and r.json()["code"] == "KEY_RECORD_IMMUTABLE", (who, p, method, r.status_code)
                assert r.json()["message"] == "关键记录不可修改，更正请追加新记录"
    assert CheckRecord.objects.get(pk=cr.id).note == "线管走向一致"

    # 3) 数据库：应用运行账号对关键表的 UPDATE / DELETE 因权限被拒（SELECT 正常）
    s = connection.settings_dict
    with psycopg.connect(dbname=s["NAME"], user="app_runtime", password="app_runtime_pw", host=s["HOST"], port=s["PORT"],
                         autocommit=True) as rt:
        rt.execute("SELECT count(*) FROM core_checkrecord").fetchone()
        for table in KEY_TABLES:
            for sql in (f"UPDATE {table} SET id = id WHERE id = -1", f"DELETE FROM {table} WHERE id = -1"):
                with pytest.raises(psycopg.errors.InsufficientPrivilege):
                    rt.execute(sql)

    # 4) 完整性校验：篡改前「完整」；超级账号（表所有者）篡改一条核对记录后报第一条不一致记录的编号与时间
    owner = client_for("owner")
    assert client_for("pm").post(f"/api/projects/{pid}/integrity-checks").status_code == 403
    r = owner.post(f"/api/projects/{pid}/integrity-checks")
    assert r.status_code == 201
    res = r.json()["result"]
    assert res["ok"] is True and res["count"] == 6 and res["tail"] == entries[-1].hash and res["firstBad"] is None
    with connection.cursor() as cur:
        cur.execute("UPDATE core_checkrecord SET note = %s WHERE id = %s", ["线管走向不一致", cr.id])
    res = client_for("admin").post(f"/api/projects/{pid}/integrity-checks").json()["result"]
    bad = LedgerEntry.objects.get(record_type="core_checkrecord", record_id=cr.id)
    assert res["ok"] is False
    assert res["firstBad"]["code"] == f"KR-{bad.seq:06d}" and res["firstBad"]["seq"] == 2
    assert res["firstBad"]["at"] == timezone.localtime(bad.created_at).isoformat(timespec="seconds")
    assert res["firstBad"]["recordLabel"] == f"检查点 {cp.code} 核对记录"
    latest = owner.get(f"/api/projects/{pid}/integrity-checks").json()["latest"]
    assert latest["result"]["ok"] is False
