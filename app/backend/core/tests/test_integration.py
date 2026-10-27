"""跨模块集成验收：模块联调后补充的断言（期望值来自契约与 AC 原文）。"""
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image

from core.models import Baseline
from core.tests.test_check import ev, make_cp  # noqa: F401 — 水电点位图 V1 已签认、二楼一项基线已签认


def pdf_bytes():
    buf = io.BytesIO()
    Image.new("RGB", (800, 560), "white").save(buf, "PDF", resolution=72)
    return buf.getvalue()


@pytest.mark.django_db
def test_ac_f04_07_resign_todos_for_pm_and_owner(project, client_for, ev):  # noqa: F811
    """AC-F04-07：新版签认后受影响的已签认基线变为「待重签」，并向项目经理与业主生成待办。"""
    wi = ev["wi"]
    dm, owner, pm = client_for("design_manager"), client_for("owner"), client_for("pm")
    r = dm.post(f"/api/projects/{project.id}/drawing-sets/{ev['ds'].id}/versions", {
        "file": SimpleUploadedFile("水电点位图-V2.pdf", pdf_bytes(), "application/pdf"),
        "changeNote": "二楼书房 B 南墙新增 2 个五孔插座", "affectedWorkItemIds": str(wi.id)}, format="multipart")
    assert r.status_code == 201, r.content
    assert owner.post(f"/api/projects/{project.id}/drawing-versions/{r.json()['id']}/sign", {}, format="json").status_code == 200
    assert Baseline.objects.get(work_item=wi).status == "待重签"
    for c in (pm, owner):
        data = c.get("/api/todos").json()
        resign = [x for x in data["items"] if x["type"] == "基线待重签"]
        assert [x["path"] for x in resign] == [f"/p/{project.id}/work-items/{wi.id}"]
        assert "水电点位图 V2" in resign[0]["due"] and "待重签" in resign[0]["sealText"]
        assert c.get("/api/todos/count").json()["total"] == data["total"]
    # 项目经理更新图纸版本重提后，待重签待办消失
    v2 = r.json()["id"]
    make_cp(pm, project, wi, point="南墙新增插座底盒标高")
    r = pm.post(f"/api/projects/{project.id}/work-items/{wi.id}/baseline/submit",
                {"drawingVersionId": v2, "cost": "1640"}, format="json")
    assert r.status_code in (200, 201), r.content
    assert not [x for x in pm.get("/api/todos").json()["items"] if x["type"] == "基线待重签"]


@pytest.mark.django_db
def test_plan_export_reimport_after_room_disabled(project, client_for):
    """停用房间后，本系统导出的计划文件可原样导回（已有工作项沿用原位置），不能把新工作项放进停用房间。"""
    import openpyxl

    from core.models import Room, WorkItem
    pm = client_for("pm")
    room = Room.objects.get(project=project, floor__name="二楼", name="书房 B")
    assert pm.post(f"/api/projects/{project.id}/rooms/{room.id}/disable", {}, format="json").status_code in (200, 201)
    resp = pm.get(f"/api/projects/{project.id}/plan/export.xlsx")
    data = b"".join(resp.streaming_content) if resp.streaming else resp.content
    r = pm.post(f"/api/projects/{project.id}/plan/imports",
                {"file": SimpleUploadedFile("原样导回.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                format="multipart")
    assert r.status_code in (200, 201), r.content
    assert r.json()["errors"] == [], r.json()["errors"][:3]
    # 新工作项放进停用房间仍被拒
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = next(w for w in wb.worksheets if w.max_row > 1 and any("名称" in str(c.value or "") or "事项" in str(c.value or "") for c in w[1]))
    head = [str(c.value or "") for c in ws[1]]
    row = [c.value for c in ws[2]]
    row[next(i for i, h in enumerate(head) if "事项" in h or "名称" in h)] = "书房 B 新增吊柜"
    for i, h in enumerate(head):
        if "房间" in h:
            row[i] = "书房 B"
        if "楼层" in h:
            row[i] = "二楼"
    ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    r = pm.post(f"/api/projects/{project.id}/plan/imports",
                {"file": SimpleUploadedFile("新增.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                format="multipart")
    assert any("书房 B" in e.get("message", str(e)) for e in r.json()["errors"]), r.json()["errors"][:3]
    assert WorkItem.objects.filter(project=project, room=room).exists()
