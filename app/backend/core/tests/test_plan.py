"""结构与计划（plan）：FLOW-04、FLOW-05、FLOW-15、FLOW-16、FLOW-19（编辑与留档）、FLOW-23。

期望值全部来自 TDD 契约里的已知业务数据并手算写死，不调用被测函数求期望值。
"""
import base64
import csv
import datetime
import io
import json

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.api import structure as st
from core.models import (AcceptancePhoto, Attendance, Baseline, BaselineSignoff, ChangeOrder, Checkpoint, CheckPhoto,
                         CheckRecord, DiscoveryCorrection, DrawingSet, DrawingSignoff, DrawingVersion, Floor,
                         ImportBatch, Issue, LedgerEntry, Notice, Phase, PlanChangeLog, ReleaseRecord, Room,
                         StoredFile, Supplement, WorkItem, WorkItemAttachment)

TODAY = datetime.date(2026, 10, 27)


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr(st, "today", lambda: TODAY)


def url(project, tail):
    return f"/api/projects/{project.id}/{tail}"


def phase(project, code):
    return Phase.objects.get(project=project, code=code)


def floor(project, code):
    return Floor.objects.get(project=project, code=code)


def room(project, code):
    return Room.objects.get(project=project, code=code)


def replace_items(project, code, rows):
    """把某工段的模板工作项停用，换成已知数据 [(名称, 权重, 完成度)]。"""
    p = phase(project, code)
    WorkItem.objects.filter(phase=p).update(is_disabled=True)
    out = []
    for name, w, pct in rows:
        wi = WorkItem.objects.create(project=project, phase=p, name=name, weight=w, progress=pct,
                                     floor=floor(project, "f2"), room=room(project, "f2-sb"))
        Baseline.objects.create(project=project, work_item=wi)
        out.append(wi)
    return out


class Evi:
    """直接用模型构造图纸、检查点、核对、照片、放行、问题与变更单。"""

    def __init__(self, project, users):
        self.p, self.u = project, users
        f = StoredFile.objects.create(project=project, storage_key="x/a.pdf", sha256="a" * 64, size=1, mime="application/pdf",
                                      original_name="建施-01.pdf")
        self.photo = StoredFile.objects.create(project=project, storage_key="x/p.jpg", sha256="b" * 64, size=1,
                                               mime="image/jpeg", original_name="现场.jpg")
        ds = DrawingSet.objects.create(project=project, name="建筑")
        self.v1 = DrawingVersion.objects.create(project=project, drawing_set=ds, version_no=1, file=f,
                                                created_by=users["design_manager"])
        DrawingSignoff.objects.create(project=project, version=self.v1, signer=users["owner"], file_sha256="a" * 64)
        self.n = 0
        self.wi = WorkItem.objects.filter(project=project, phase__code="p1").first()

    def cp(self, concealed, status="未核对"):
        self.n += 1
        return Checkpoint.objects.create(project=self.p, work_item=self.wi, code=f"CK-{self.n:03d}", point="门洞净宽 900mm",
                                         is_concealed=concealed, drawing_version=self.v1, x=0.5, y=0.5, status=status,
                                         created_by=self.u["qc"])

    def check(self, cp, conclusion="一致", latest=True, photo=True, phase_="", covered=False):
        cr = CheckRecord.objects.create(project=self.p, checkpoint=cp, checker=self.u["qc"], drawing_version=self.v1,
                                        was_latest_signed=latest, conclusion=conclusion, was_covered=covered,
                                        discovery_phase=phase_)
        if photo:
            CheckPhoto.objects.create(project=self.p, check_record=cr, file=self.photo, file_sha256="b" * 64,
                                      pin_x=0.5, pin_y=0.5)
        return cr

    def release(self, cp, cr):
        r = ReleaseRecord.objects.create(project=self.p, checkpoint=cp, releaser=self.u["owner"], is_proxy=False,
                                         basis_check=cr)
        cp.status, cp.released_at = "已放行", r.released_at
        cp.save()
        return r

    def issue(self, cr, status="未闭环", due=None, seq=None):
        return Issue.objects.create(project=self.p, seq_no=seq or Issue.objects.filter(project=self.p).count() + 1,
                                    source_check=cr, checkpoint=cr.checkpoint, work_item=self.wi, phase=self.wi.phase,
                                    floor=self.wi.floor, room=self.wi.room, description="门洞现场净宽 860mm，与图纸 900mm 不符",
                                    level="重要", found_on=TODAY, due_date=due, status=status)


# ------------------------------------------------------------------ FLOW-04 结构

def test_flow04_structure_crud_and_disable(client_for, project, users):
    pm = client_for("pm")
    f2 = floor(project, "f2")
    r = pm.post(url(project, "rooms"), {"name": "茶室", "floorId": f2.id}, format="json")
    assert r.status_code == 201
    rid = r.json()["id"]
    assert pm.patch(url(project, f"rooms/{rid}"), {"name": "品茶室"}, format="json").json()["name"] == "品茶室"
    assert pm.post(url(project, "rooms"), {"name": "品茶室", "floorId": f2.id}, format="json").json()["fields"][0]["message"] == "已存在同名房间"
    rows = pm.get(url(project, "floor-view")).json()["roomMatrix"]
    assert any(x["name"] == "品茶室" for x in rows)

    # 无记录的楼层可删除，并从矩阵与筛选项中消失
    fid = pm.post(url(project, "floors"), {"name": "阁楼"}, format="json").json()["id"]
    assert any(x["name"] == "阁楼" for x in pm.get(url(project, "floor-view")).json()["floorMatrix"])
    assert pm.delete(url(project, f"floors/{fid}")).status_code == 200
    assert not any(x["name"] == "阁楼" for x in pm.get(url(project, "floor-view")).json()["floorMatrix"])

    # 有核对记录的房间：删除被拒 → 停用 → 不再出现在新增选项，历史仍可查
    e = Evi(project, users)
    sb = room(project, "f2-sb")
    e.wi = WorkItem.objects.filter(project=project, room=sb).first()
    e.check(e.cp(True))
    r = pm.delete(url(project, f"rooms/{sb.id}"))
    assert r.status_code == 409 and r.json()["code"] == "HAS_RECORDS_DISABLE_ONLY"
    assert r.json()["message"] == "已有记录，只能停用"
    assert Room.objects.filter(pk=sb.id).exists()
    assert pm.post(url(project, f"rooms/{sb.id}/disable")).json()["isDisabled"] is True
    s = pm.get(url(project, "structure")).json()
    assert next(x for x in s["rooms"] if x["id"] == sb.id)["isDisabled"] is True
    assert not any(x["name"] == "书房 B" for x in pm.get(url(project, "floor-view")).json()["roomMatrix"])
    r = pm.post(url(project, "work-items"), {"name": "书房 B 补墙", "phaseId": phase(project, "p1").id,
                                            "floorId": f2.id, "roomId": sb.id, "weight": 1, "progress": 0}, format="json")
    assert r.status_code == 400 and r.json()["fields"][0]["message"] == "该房间已停用，请重新选择"
    assert pm.get(url(project, f"scope-detail?room={sb.id}")).status_code == 200

    # 有下级房间的楼层只能停用；只读角色不能改结构
    r = pm.delete(url(project, f"floors/{f2.id}"))
    assert r.status_code == 409 and r.json()["details"]["children"] > 0
    r = client_for("owner").post(url(project, "floors"), {"name": "地下二层"}, format="json")
    assert r.status_code == 403


def test_flow04_work_item_rules_and_weighted_progress(client_for, project):
    pm = client_for("pm")
    p1 = phase(project, "p1")
    WorkItem.objects.filter(phase=p1).update(is_disabled=True)
    f2, f3 = floor(project, "f2"), floor(project, "f3")

    r = pm.post(url(project, "work-items"), {}, format="json")
    assert r.status_code == 400
    fields = {f["field"]: f["message"] for f in r.json()["fields"]}
    assert fields["name"] == "请填写工作项名称" and fields["phaseId"] == "请选择所属工段"
    r = pm.post(url(project, "work-items"), {"name": "测试", "phaseId": p1.id, "weight": 7, "progress": 120}, format="json")
    fields = {f["field"]: f["message"] for f in r.json()["fields"]}
    assert fields == {"weight": "权重为 1–5 的整数", "progress": "完成度为 0–100 的整数"}
    # 房间不属于所选楼层（AC-F03-04）
    r = pm.post(url(project, "work-items"), {"name": "三楼书房 B 隔墙", "phaseId": p1.id, "floorId": f3.id,
                                            "roomId": room(project, "f2-sb").id, "weight": 2, "progress": 0}, format="json")
    assert r.status_code == 400 and r.json()["message"] == "房间与楼层不匹配"

    # 权重 2、2、1，完成度 100、50、0 → (2×100+2×50+1×0)÷5 = 60.0%，进行中（AC-F03-05、AC-F03-06）
    for name, w, pct in (("二楼书房 B 隔墙砌筑", 2, 100), ("二楼卧室 B 隔墙砌筑", 2, 50), ("二楼卫生间 B 隔墙", 1, 0)):
        r = pm.post(url(project, "work-items"), {"name": name, "phaseId": p1.id, "floorId": f2.id, "weight": w,
                                                "progress": pct, "description": "按图砌筑", "acceptanceCriteria": "垂直度 ≤5mm"},
                    format="json")
        assert r.status_code == 201
    ph = next(x for x in pm.get(url(project, "structure")).json()["phases"] if x["id"] == p1.id)
    assert ph["progress"] == 60.0 and ph["status"] == "进行中"
    assert pm.get(url(project, f"phases/{p1.id}/detail")).json()["phase"]["progress"] == 60.0

    # 工序规则：名称命中「管道井」规则补齐，已填字段不被覆盖（AC-F03-10）
    r = pm.post(url(project, "work-items"), {"name": "二楼卫生间 A 隔墙及管道井", "phaseId": p1.id, "floorId": f2.id,
                                            "roomId": room(project, "f2-wca").id, "weight": 1, "progress": 0}, format="json")
    assert r.json()["description"].startswith("隔墙采用 120mm 加气混凝土砌块，管道井随砌随封")
    assert r.json()["acceptanceCriteria"] != ""
    r = pm.post(url(project, "work-items"), {"name": "二楼卫生间 B 隔墙及管道井", "phaseId": p1.id, "floorId": f2.id,
                                            "description": "业主要求改用轻质砖", "weight": 1, "progress": 0}, format="json")
    assert r.json()["description"] == "业主要求改用轻质砖" and r.json()["acceptanceCriteria"] != ""
    assert r.json()["ownerName"] == "待指派" and r.json()["baselineStatus"] == "草稿"


def test_work_item_baseline_lock_delete_disable(client_for, project, users):
    pm = client_for("pm")
    wi = WorkItem.objects.filter(project=project, phase__code="p1").first()
    bl = wi.baseline
    bl.status = "已签认"
    bl.save()
    r = pm.patch(url(project, f"work-items/{wi.id}"), {"description": "改说明", "progress": 40,
                                                      "actualFinish": "2026-10-20"}, format="json")
    assert r.status_code == 409 and r.json()["code"] == "BASELINE_LOCKED"
    assert r.json()["message"] == "已签认，请发起变更单或等待新版图纸重签"
    wi.refresh_from_db()
    assert wi.progress == 40 and wi.actual_finish == datetime.date(2026, 10, 20) and wi.description != "改说明"

    bl.status = "待签认"
    bl.save()
    r = pm.patch(url(project, f"work-items/{wi.id}"), {"name": "改名", "progress": 60}, format="json")
    assert r.status_code == 409 and r.json()["message"] == "基线待业主签认，需业主退回后才能修改"
    wi.refresh_from_db()
    assert wi.progress == 60 and wi.name != "改名"
    assert pm.patch(url(project, f"work-items/{wi.id}"), {"progress": 70}, format="json").status_code == 200

    bl.status = "草稿"
    bl.save()
    r = pm.patch(url(project, f"work-items/{wi.id}"), {"floorId": floor(project, "f3").id}, format="json")
    if wi.room_id:
        assert r.status_code == 400 and r.json()["message"] == "房间与楼层不匹配"
    assert client_for("owner").patch(url(project, f"work-items/{wi.id}"), {"progress": 1}, format="json").status_code == 403

    # 有关联记录的工作项只能停用；无记录的可删除
    e = Evi(project, users)
    e.wi = wi
    e.check(e.cp(True))
    r = pm.delete(url(project, f"work-items/{wi.id}"))
    assert r.status_code == 409 and r.json()["message"] == "已有记录，只能停用"
    assert pm.post(url(project, f"work-items/{wi.id}/disable")).json()["isDisabled"] is True
    ids = [x["id"] for x in pm.get(url(project, "work-items?phase=%d" % wi.phase_id)).json()["items"]]
    assert wi.id not in ids
    other = WorkItem.objects.filter(project=project, phase__code="p6").first()
    assert pm.delete(url(project, f"work-items/{other.id}")).status_code == 200
    assert not WorkItem.objects.filter(pk=other.id).exists()


# ------------------------------------------------------------------ FLOW-19 编辑工段计划并留档

def test_flow19_phase_plan_log(client_for, project):
    pm = client_for("pm")
    p1 = phase(project, "p1")
    body = {"plannedStart": "2026-09-19", "plannedEnd": "2026-11-08", "teamLead": p1.team_lead, "weight": p1.weight,
            "scopeNote": p1.scope_note, "reason": "书房 B 门洞待设计变更，相关隔墙暂缓"}
    r = pm.patch(url(project, f"phases/{p1.id}/plan"), body, format="json")
    assert r.status_code == 200 and r.json()["changed"] is True
    logs = list(PlanChangeLog.objects.filter(phase=p1))
    assert len(logs) == 1
    log = logs[0]
    assert [(c["field"], c["before"], c["after"]) for c in log.changes] == [("plannedEnd", "2026-10-18", "2026-11-08")]
    assert log.duration_delta_days == 21 and log.reason == "书房 B 门洞待设计变更，相关隔墙暂缓" and log.at is not None
    assert LedgerEntry.objects.filter(record_type="core_planchangelog", record_id=log.id).exists()  # 关键记录入链

    r = pm.patch(url(project, f"phases/{p1.id}/plan"), body, format="json")
    assert r.json()["changed"] is False and PlanChangeLog.objects.filter(phase=p1).count() == 1

    r = pm.patch(url(project, f"phases/{p1.id}/plan"), {**body, "plannedStart": "2026-11-09"}, format="json")
    assert r.status_code == 400 and r.json()["message"] == "开始日期不能晚于结束日期"
    assert PlanChangeLog.objects.filter(phase=p1).count() == 1
    assert client_for("owner").patch(url(project, f"phases/{p1.id}/plan"), body, format="json").status_code == 403
    detail = pm.get(url(project, f"phases/{p1.id}/detail")).json()
    assert detail["planLogs"][0]["durationDeltaDays"] == 21 and detail["phase"]["plannedEnd"] == "2026-11-08"


# ------------------------------------------------------------------ FLOW-15 总览与四项指标

def test_flow15_overview_kpi_and_evidence(client_for, project, users):
    replace_items(project, "p1", [("二楼书房 B 隔墙砌筑", 2, 100), ("二楼卧室 B 隔墙砌筑", 2, 50), ("二楼卫生间 B 隔墙", 1, 0)])
    replace_items(project, "p2", [("二楼书房 B 强电布管", 1, 10)])
    e = Evi(project, users)
    now = timezone.now()
    # 隐蔽检查点 8 个已放行或已覆盖，6 个合规
    for k in range(8):
        cp = e.cp(True)
        cr = e.check(cp, photo=(k != 7))           # 第 8 个：核对没有照片
        e.release(cp, cr)
        if k == 6:                                  # 第 7 个：放行前已被覆盖（未放行先覆盖）
            cp.covered_at, cp.unreleased_cover = now - datetime.timedelta(days=1), True
            cp.save()
    # 不一致 11 条：1 条问题已作废；其余 10 条中 7 条覆盖前（其中 1 条由覆盖后修正为覆盖前）
    phases_ = ["覆盖前"] * 6 + ["覆盖后"] * 4
    for k, ph in enumerate(phases_):
        cp = e.cp(False, "不一致")
        cr = e.check(cp, "不一致", latest=(k < 8), phase_=ph)
        e.issue(cr, due=datetime.date(2026, 10, 20) if k == 0 else datetime.date(2026, 11, 3))
        if k == 6:
            DiscoveryCorrection.objects.create(project=project, check_record=cr, before="覆盖后", after="覆盖前",
                                               reason="质检复核：当时墙体尚未抹灰", actor=users["qc"])
    void_cr = e.check(e.cp(False, "不一致"), "不一致", phase_="覆盖后")
    e.issue(void_cr, status="已作废")
    e.check(e.cp(False, "待放行"))                 # 第 20 条核对
    # 已生效变更单 4 张，其中 1 张后补；另有 1 张草稿不计
    for k, (status, late) in enumerate([("已生效", False), ("已生效", False), ("已生效", False), ("已生效", True),
                                        ("草稿", None)]):
        co = ChangeOrder.objects.create(project=project, seq_no=k + 1, status=status, built_before_sign=late,
                                        created_by=users["pm"])
        co.work_items.set([e.wi])
    assert CheckRecord.objects.filter(project=project).count() == 20

    ov = client_for("owner").get(url(project, "overview"))
    assert ov.status_code == 200
    o = ov.json()
    k = o["kpi"]
    # 总体 = (8×60 + 22×10 + 0)÷100 = 7.0；计划基准 = (8×100 + 22×8÷91×100)÷100 = 9.934 → 9.9；滞后 2.9
    assert (k["overall"], k["planned"], k["gap"]) == (7.0, 9.9, -2.9)
    assert (k["periodStart"], k["periodEnd"], k["duration"], k["used"]) == ("2026-09-19", "2027-08-30", 346, 39)
    assert k["todayAttendance"] is None
    assert [m["value"] for m in o["evidence"]] == [70.0, 75.0, 90.0, 75.0]
    assert [(m["numerator"], m["denominator"]) for m in o["evidence"]] == [(7, 10), (6, 8), (18, 20), (3, 4)]
    assert [m["target"] for m in o["evidence"]] == [90, 100, 100, 100]
    assert k["issues"] == 10 and k["legacy"] == 10 and k["overdueOpen"] == 1
    assert o["legacy"][0]["overdueDays"] == 7
    p1 = next(p for p in o["phases"] if p["name"] == "砌墙")
    assert (p1["progress"], p1["planned"], p1["status"]) == (60.0, 100.0, "进行中")
    # 楼层矩阵：二楼 × 砌墙 = 60.0，全楼通用 × 砌墙无事项 → null
    row = next(x for x in o["floorMatrix"] if x["name"] == "二楼")
    assert row["cells"][0] == 60.0
    assert next(x for x in o["floorMatrix"] if x["id"] == "all")["cells"][0] is None
    # 按工段筛选：无数据的工段四项均为「暂无数据」（null）
    tile = phase(project, "p3")
    ev = client_for("owner").get(url(project, f"overview?phaseId={tile.id}")).json()["evidence"]
    assert [m["value"] for m in ev] == [None, None, None, None]
    ev = client_for("owner").get(url(project, f"overview?phaseId={phase(project, 'p1').id}")).json()["evidence"]
    assert [m["value"] for m in ev] == [70.0, 75.0, 90.0, 75.0]


def test_flow15_floor_view_and_drilldown(client_for, project):
    replace_items(project, "p1", [("二楼书房 B 隔墙砌筑", 2, 100), ("二楼书房 B 门洞过梁", 2, 50), ("二楼书房 B 塞缝", 1, 0)])
    c = client_for("owner")
    fv = c.get(url(project, "floor-view")).json()
    sb = room(project, "f2-sb")
    row = next(x for x in fv["roomMatrix"] if x["id"] == sb.id)
    assert row["cells"][0] == 60.0 and row["cells"][1] is not None  # 水电模板中书房 B 有事项（完成度 0）
    d = c.get(url(project, f"scope-detail?room={sb.id}")).json()
    assert d["title"] == "二楼 · 书房 B" and d["groups"][0]["name"] == "砌墙" and len(d["groups"][0]["items"]) == 3
    detail = c.get(url(project, f"phases/{phase(project, 'p1').id}/detail")).json()
    assert len([x for x in detail["items"] if x["roomId"] == sb.id]) == 3


# ------------------------------------------------------------------ FLOW-16 出勤

def test_flow16_attendance(client_for, project):
    pm = client_for("pm")
    p1, p2 = phase(project, "p1"), phase(project, "p2")
    pm.patch(url(project, f"phases/{p1.id}/plan"), {"plannedStart": "2026-09-19", "plannedEnd": "2026-11-08"}, format="json")
    r = pm.put(url(project, f"phases/{p1.id}/attendance/2026-10-20"), {"headcount": 6}, format="json")
    assert r.status_code == 200
    a = pm.get(url(project, f"phases/{p1.id}/detail")).json()["attendanceStats"]
    assert (a["total"], a["peak"], a["avg"], a["workDays"]) == (6, 6, 6.0, 1)
    r = pm.post(url(project, f"phases/{p1.id}/attendance/batch"),
                {"start": "2026-10-21", "end": "2026-10-27", "headcount": 5, "mode": "workdays"}, format="json")
    assert r.json()["written"] == 6  # 10-25 为周日
    assert not Attendance.objects.filter(phase=p1, date=datetime.date(2026, 10, 25)).exists()
    r = pm.put(url(project, f"phases/{p1.id}/attendance/2026-10-26"), {"headcount": 41}, format="json")
    assert r.status_code == 400 and r.json()["fields"][0]["message"] == "请填写 0–40 的整数"
    r = pm.put(url(project, f"phases/{p1.id}/attendance/2026-12-01"), {"headcount": 3}, format="json")
    assert r.status_code == 400 and r.json()["fields"][0]["field"] == "date"
    assert client_for("owner").put(url(project, f"phases/{p1.id}/attendance/2026-10-26"), {"headcount": 3},
                                   format="json").status_code == 403
    pm.put(url(project, f"phases/{p2.id}/attendance/2026-10-21"), {"headcount": 3}, format="json")

    m = pm.get(url(project, "attendance?month=2026-10")).json()
    assert m["days"]["2026-10-21"] == 8 and "2026-10-19" not in m["days"]  # 未登记不推算
    assert (m["total"], m["workDays"], m["peak"], m["peakDate"]) == (39, 7, 8, "2026-10-21")
    assert m["unregistered"] == 20  # 10-01 ~ 10-27 共 27 天，登记 7 天
    a = pm.get(url(project, f"phases/{p1.id}/detail")).json()["attendanceStats"]
    assert (a["total"], a["peak"], a["workDays"]) == (36, 6, 7)
    assert pm.get(url(project, "overview")).json()["kpi"]["todayAttendance"] == 5

    r = pm.post(url(project, f"phases/{p1.id}/attendance/batch"),
                {"start": "2026-10-26", "end": "2026-11-02", "headcount": 4, "mode": "untilToday"}, format="json")
    assert r.json()["written"] == 2 and r.json()["overwritten"] == 2
    r = pm.post(url(project, f"phases/{p1.id}/attendance/batch"),
                {"start": "2026-11-01", "end": "2026-11-01", "headcount": 4, "mode": "workdays"}, format="json")
    assert r.status_code == 400  # 11-01 是周日，没有可写日期


def _png():
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 120, 40)).save(b, "PNG")
    return b.getvalue()


def test_acceptance_photo_cycle(client_for, project):
    qc = client_for("qc")
    p1 = phase(project, "p1")
    f = qc.post(url(project, "files"), {"file": SimpleUploadedFile("砌体验收-二楼.png", _png(), "image/png")},
                format="multipart").json()
    r = qc.post(url(project, f"phases/{p1.id}/acceptance-photos"), {"fileIds": [f["id"]]}, format="json")
    assert r.status_code == 201
    ph = r.json()["items"][0]
    assert ph["status"] == "待复验" and ph["caption"] == "砌体验收-二楼"
    seq = [qc.post(url(project, f"acceptance-photos/{ph['id']}/status")).json()["status"] for _ in range(3)]
    assert seq == ["已验收", "施工中", "待复验"]
    assert client_for("owner_member").post(url(project, f"acceptance-photos/{ph['id']}/status")).status_code == 403


# ------------------------------------------------------------------ FLOW-05 Excel / CSV

TASK_HEAD = ["序号", "施工阶段", "工作事项名称", "工作说明", "验收标准", "所属楼层", "所属房间 / 区域", "计划完成日期",
             "实际完成日期", "负责人", "权重(1-5)", "完成度(%)"]


def xlsx(rows, head=TASK_HEAD):
    wb = Workbook()
    ws = wb.active
    ws.title = "工作事项"
    ws.append(head)
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_flow05_template_export_formats(client_for, project):
    pm = client_for("pm")
    for tail in ("plan/template.xlsx", "plan/export.xlsx"):
        r = pm.get(url(project, tail))
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["施工阶段", "工作事项", "补充事项", "填写说明", "取值"]
        ps = wb["施工阶段"]
        assert ps["A2"].data_type == "n" and ps["C2"].data_type == "s" and ps["C2"].number_format == "@"
        assert ps["F2"].data_type == "n" and ps["C2"].value == "2026-09-19"
    ws = load_workbook(io.BytesIO(pm.get(url(project, "plan/export.xlsx")).content))["工作事项"]
    assert ws.max_row == 188  # 表头 + 187 项
    head = [c.value for c in ws[1]]
    assert head == TASK_HEAD
    row = [c for c in ws[3]]
    assert row[1].value == "砌墙" and row[2].value == "负一楼车库隔墙砌筑"
    assert row[7].data_type == "s" and row[7].value == "2026-09-20" and row[7].number_format == "@"
    assert row[10].data_type == "n" and row[10].value == 2 and row[11].data_type == "n"
    assert client_for("owner").get(url(project, "plan/export.xlsx")).status_code == 403


def test_flow05_import_validate_and_execute(client_for, project):
    pm = client_for("pm")
    bad = xlsx([
        [1, "水电安装", "负一楼新增配电箱", "", "", "负一楼", "车库", "2026-10-30", "", "赵师傅", 2, 0],
        [2, "砌墙", "三楼书房 B 隔墙", "", "", "三楼", "书房 B", "2026-10-30", "", "陈师傅", 2, 0],
        [3, "砌墙", "一楼餐厅新砌隔墙", "", "", "一楼", "餐厅", "11月31日", "", "陈师傅", 2, 0],
        [4, "砌墙", "一楼厨房新砌隔墙", "", "", "一楼", "厨房", "2026-09-28", "", "陈师傅", 7, 0],
        [5, "砌墙", "一楼读书区隔墙砌筑", "", "", "一楼", "读书区", "2026-09-28", "", "陈师傅", 2, 120],
    ])
    r = pm.post(url(project, "plan/imports"), {"file": SimpleUploadedFile("君华别墅-施工计划-1027修订.xlsx", bad)},
                format="multipart")
    assert r.status_code == 201
    errs = r.json()["errors"]
    assert [(x["sheet"], x["row"], x["column"]) for x in errs] == [
        ("工作事项", 2, "施工阶段"), ("工作事项", 3, "所属房间 / 区域"), ("工作事项", 4, "计划完成日期"),
        ("工作事项", 5, "权重(1-5)"), ("工作事项", 6, "完成度(%)")]
    assert errs[0]["message"].startswith("阶段名「水电安装」不可识别")
    assert "不属于楼层「三楼」" in errs[1]["message"] and errs[2]["message"] == "日期「11月31日」无法识别"
    assert errs[3]["message"] == "权重「7」越界，应为 1–5 的整数"
    r = pm.post(url(project, f"plan/imports/{r.json()['id']}/execute"))
    assert r.status_code == 409 and r.json()["code"] == "IMPORT_HAS_ERRORS"
    assert pm.post(url(project, "plan/imports"), {"file": SimpleUploadedFile("计划.pdf", b"%PDF")},
                   format="multipart").status_code == 415

    signed = WorkItem.objects.get(project=project, name="负一楼车库隔墙砌筑")
    signed.baseline.status = "已签认"
    signed.baseline.save()
    before = {f: getattr(signed, f) for f in ("planned_finish", "owner_name", "weight", "progress", "description")}
    good = xlsx([
        [1, "砌墙", "负一楼车库隔墙砌筑", "", "", "负一楼", "车库", "09-22", "", "陈师傅", 2, 100],
        [2, "砌墙", "负一楼保姆间隔墙砌筑", "", "", "负一楼", "保姆间", "2026-09-21", "", "刘师傅", 2, 100],
        [3, "砌墙", "负一楼影音室隔墙砌筑", "", "", "负一楼", "", "2026-09-24", "", "陈师傅", 2, 0],
    ])
    r = pm.post(url(project, "plan/imports"), {"file": SimpleUploadedFile("君华别墅-施工计划-1027修订.xlsx", good)},
                format="multipart")
    b = r.json()
    assert b["errors"] == [] and b["preview"]["add"] == 1 and b["preview"]["update"] == 1
    assert [x["name"] for x in b["preview"]["signed"]] == ["负一楼车库隔墙砌筑"]
    r = pm.post(url(project, f"plan/imports/{b['id']}/execute"))
    assert r.status_code == 200
    res = r.json()["result"]
    assert (res["added"], res["updated"]) == (1, 1)
    assert res["signed"][0]["note"] == "已签认，需走变更单" and "计划日期 09-20 → 09-22" in res["signed"][0]["changes"]
    signed.refresh_from_db()
    assert {f: getattr(signed, f) for f in before} == before  # 已签认项字段完全不变
    nanny = WorkItem.objects.get(project=project, name="负一楼保姆间隔墙砌筑")
    assert (nanny.owner_name, nanny.progress) == ("刘师傅", 100)
    new = WorkItem.objects.get(project=project, name="负一楼影音室隔墙砌筑")
    assert new.source == "excel_import" and new.floor.name == "负一楼" and new.room is None
    assert new.description != ""  # 工序规则补齐（隔墙砌筑）
    assert pm.post(url(project, f"plan/imports/{b['id']}/execute")).status_code == 409


def test_flow05_csv_gbk(client_for, project):
    pm = client_for("pm")
    s = io.StringIO()
    w = csv.writer(s)
    w.writerow(TASK_HEAD)
    w.writerow([1, "水电", "二楼书房 B 强电箱移位", "", "", "二楼", "书房 B", "2026-11-05", "", "赵师傅", 3, 0])
    r = pm.post(url(project, "plan/imports"),
                {"file": SimpleUploadedFile("水电补充.csv", s.getvalue().encode("gbk"))}, format="multipart")
    assert r.json()["errors"] == [] and r.json()["preview"]["add"] == 1
    assert pm.post(url(project, f"plan/imports/{r.json()['id']}/execute")).json()["result"]["added"] == 1
    wi = WorkItem.objects.get(project=project, name="二楼书房 B 强电箱移位")
    assert (wi.phase.name, wi.room.name, wi.planned_finish, wi.weight) == ("水电", "书房 B", datetime.date(2026, 11, 5), 3)


# ------------------------------------------------------------------ FLOW-23 原型 JSON 导入

def backup():
    png = "data:image/png;base64," + base64.b64encode(_png()).decode()
    return {
        "phase": {"p1": {"start": "2026-09-19", "end": "2026-10-18", "crew": "瓦工班 · 陈师傅（6人）", "w": 0.08}},
        "tasks": {"p1": [
            {"n": "放线定位、材料进场验收", "note": "", "std": "", "plan": "09-19", "act": "09-19", "owner": "陈师傅",
             "fl": "all", "rm": "all", "w": 1, "pct": 100, "files": ["F1"]},
            {"n": "负一楼卫生间新砌隔墙及管道井封堵", "plan": "09-22", "owner": "陈师傅", "fl": "b1", "rm": "b1-wc", "w": 2,
             "pct": 100, "files": ["F2"]},
        ]},
        "att": {"p1": {"2026-09-19": 8, "2026-09-20": 6}},
        "issues": [{"id": "N123456", "pid": "p1", "t": "负一楼车库门洞偏位 30mm", "d": "2026-09-25", "lv": "重要",
                    "fix": "拆除重砌", "owner": "陈师傅", "due": "2026-10-02", "fl": "b1", "rm": "b1-car", "st": "未闭环"}],
        "edit": {}, "delIssues": {}, "closed": {"N123456": "2026-10-01"}, "photos": {"p1": [
            {"cap": "负一楼卫生间管道井封堵", "st": "已验收", "src": png, "date": "2026-09-23"}]},
        "upto": "", "log": [{"t": "2026-09-30 18:20", "pid": "p1", "pname": "砌墙", "note": "材料晚到",
                             "items": [{"f": "end", "from": "2026-10-15", "to": "2026-10-18", "dd": 3}]}],
        "files": [{"id": "F1", "sec": "task", "fo": "", "name": "砌块出厂检测报告-MU10.pdf", "type": "application/pdf",
                   "size": 204800},
                  {"id": "F2", "sec": "task", "fo": "", "name": "负一楼卫生间管道井封堵-完成面.jpg", "type": "image/jpeg",
                   "size": 1024},
                  {"id": "F3", "sec": "ct", "fo": "c1", "name": "开工交底记录-20260919.pdf", "type": "application/pdf",
                   "size": 4096}],
        "folders": {"ct": [{"id": "c1", "name": "开工资料"}]},
        "extra": [{"id": "E1", "title": "业主要求增加车库储物架", "desc": "口头约定", "date": "2026-09-28",
                   "due": "2026-10-10", "st": "处理中", "owner": "张工", "files": []}],
        "notice": [{"id": "M1", "board": "owner", "who": "刘先生", "text": "周六上午来看现场", "ts": "2026-10-08 09:00", "re": ""},
                   {"id": "M2", "board": "owner", "who": "张工", "text": "好的，九点在门口等", "ts": "2026-10-08 09:30",
                    "re": "M1"}],
        "noticeName": {},
    }


def test_flow23_prototype_import(client_for, project):
    admin = client_for("admin")
    up = lambda c, data, name="君华别墅项目数据-2026-10-09.json": c.post(  # noqa: E731
        url(project, "prototype-imports"), {"file": SimpleUploadedFile(name, data)}, format="multipart")
    r = up(admin, json.dumps(backup()).encode())
    assert r.status_code == 409 and "原型备份只能导入「筹备」状态的项目" in r.json()["message"]
    project.status = "筹备"
    project.save()
    assert up(client_for("pm"), json.dumps(backup()).encode()).status_code == 403

    r = up(admin, json.dumps(backup(), ensure_ascii=False).encode("utf-8"))
    assert r.status_code == 201, r.json()
    out = r.json()
    assert out["status"] == "已导入"
    assert out["counts"] == {"workItems": 2, "issues": 1, "attendance": 2, "planLogs": 1, "supplements": 1,
                             "notices": 2, "photos": 1}
    assert {x["name"] for x in out["missing"]} == {"开工交底记录-20260919.pdf", "负一楼卫生间管道井封堵-完成面.jpg",
                                                   "砌块出厂检测报告-MU10.pdf"}
    p1 = phase(project, "p1")
    assert WorkItem.objects.filter(phase=p1, is_disabled=False).count() == 2  # 整段替换
    assert set(WorkItem.objects.filter(phase=p1).values_list("source", flat=True)) == {"prototype_import"}
    assert Attendance.objects.get(phase=p1, date=datetime.date(2026, 9, 19)).source == "prototype_import"
    iss = Issue.objects.get(project=project, source="prototype_import")
    assert (iss.status, iss.closed_on, iss.room.name) == ("已闭环", datetime.date(2026, 10, 1), "车库")
    log = PlanChangeLog.objects.get(project=project)
    assert (log.source, log.duration_delta_days, log.changes[0]["originalAt"]) == ("prototype_import", 3, "2026-09-30 18:20")
    assert Supplement.objects.get(project=project).source == "prototype_import"
    assert Notice.objects.filter(project=project, source="prototype_import").count() == 2
    assert Notice.objects.get(body="好的，九点在门口等").parent.body == "周六上午来看现场"
    photo = AcceptancePhoto.objects.get(project=project)
    assert (photo.source, photo.status, photo.file.status) == ("prototype_import", "已验收", "normal")
    assert StoredFile.objects.filter(project=project, status="missing").count() == 3
    assert WorkItemAttachment.objects.filter(project=project, file__status="missing").count() == 2
    assert p1.weight == 8
    # 不生成任何签认、核对或放行记录
    assert BaselineSignoff.objects.count() == 0 and CheckRecord.objects.count() == 0 and ReleaseRecord.objects.count() == 0
    assert DrawingSignoff.objects.count() == 0

    # 无效文件被拒，项目数据不变（含导入过程中才发现的错误 → 整单回滚）
    snap = (WorkItem.objects.filter(project=project).count(), Attendance.objects.filter(project=project).count(),
            Issue.objects.filter(project=project).count())
    r = up(admin, json.dumps({"name": "君华别墅备份", "rows": []}).encode(), "君华别墅备份.json")
    assert r.status_code == 422 and r.json()["message"] == "不是有效的原型备份文件"
    broken = backup()
    broken["att"] = {"p2": {"2026-10-20": 5}}
    broken["log"].append({"t": "2026-10-01 10:00", "pid": "p9", "items": []})
    r = up(admin, json.dumps(broken).encode())
    assert r.status_code == 422
    assert r.json()["code"] == "INVALID_PROTOTYPE_BACKUP"
    assert snap == (WorkItem.objects.filter(project=project).count(), Attendance.objects.filter(project=project).count(),
                    Issue.objects.filter(project=project).count())
    assert up(admin, b"\x00not json", "君华别墅备份.xlsx").status_code == 422
    hist = admin.get(url(project, "prototype-imports")).json()["items"]
    assert [h["status"] for h in hist] == ["已拒绝", "已拒绝", "已拒绝", "已导入"]
    assert ImportBatch.objects.filter(project=project, kind="prototype_json").count() == 4
