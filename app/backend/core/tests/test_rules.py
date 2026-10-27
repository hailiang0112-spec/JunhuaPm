"""横切规则验收（TDD 契约 RULE-01 / 02 / 03 / 07）。期望值取自契约中的已知业务数据。"""
import json
import logging

import pytest
from rest_framework.test import APIClient

from core.models import AuditLog, User


def _invite(client_for, project, phone, roles):
    r = client_for("pm").post(f"/api/projects/{project.id}/invitations", {"phone": phone, "roles": roles}, format="json")
    assert r.status_code == 201, r.content
    return r.json()["link"].rsplit("/", 1)[1]


@pytest.mark.django_db
def test_rule03_password_hash_rules_and_temp_password_not_logged(project, client_for, users, caplog):
    caplog.set_level(logging.DEBUG)
    token = _invite(client_for, project, "13912345678", ["qc"])
    c = APIClient()
    r = c.post(f"/api/invitations/{token}/accept", {"phone": "13912345678", "name": "孙质检", "password": "Jh2026abc"},
               format="json")
    assert r.status_code == 200, r.content
    u = User.objects.get(phone="13912345678")
    # 只存加盐哈希（argon2），不存明文
    assert u.password.startswith("argon2$") and "Jh2026abc" not in u.password
    # 纯数字、纯字母密码被拒
    for bad in ("12345678", "abcdefgh"):
        r = c.post("/api/me/password", {"oldPassword": "Jh2026abc", "newPassword": bad}, format="json")
        assert r.status_code == 400
        assert {"field": "newPassword", "message": "密码至少 8 位，且必须同时包含字母和数字"} in r.json()["fields"]
    assert User.objects.get(pk=u.pk).check_password("Jh2026abc")

    # 管理员重置：临时密码只在本次响应中出现一次，不落日志、不出现在其他接口
    admin = client_for("admin")
    r = admin.post(f"/api/admin/users/{u.id}/reset-password")
    assert r.status_code == 200 and r["Cache-Control"] == "no-store"
    temp = r.json()["tempPassword"]
    assert User.objects.get(pk=u.pk).check_password(temp)
    assert temp not in caplog.text
    for log in AuditLog.objects.all():
        assert temp not in json.dumps([log.object_label, log.before, log.after], ensure_ascii=False, default=str)
    for path in ("/api/me", f"/api/projects/{project.id}/members"):
        assert temp not in admin.get(path).content.decode()
    c2 = APIClient()
    r = c2.post("/api/auth/login", {"phone": "13912345678", "password": temp}, format="json")
    assert r.status_code == 200 and r.json()["mustChangePassword"] is True
    assert temp not in r.content.decode()


# ---------------------------------------------------------------- RULE-01 越权矩阵
def _project_routes():
    """遍历全部带 <int:project_id> 的路由及其方法，构成逐接口越权矩阵。"""
    from django.urls import get_resolver
    from django.urls.resolvers import RoutePattern

    out = []

    def walk(r, prefix):
        for x in r.url_patterns:
            if hasattr(x, "url_patterns"):
                walk(x, prefix + str(x.pattern))
            elif isinstance(x.pattern, RoutePattern) and "<int:project_id>" in str(x.pattern):
                v = x.callback.view_class
                for m in ("get", "post", "put", "patch", "delete"):
                    if hasattr(v, m) and m in getattr(v, "http_method_names", [m]):
                        out.append((m, prefix + str(x.pattern)))

    walk(get_resolver(), "/")
    return out


@pytest.mark.django_db
def test_rule01_cross_project_matrix(project, company, users, client_for):
    import datetime
    import re

    from core.models import Template
    from core.services import storage
    from core.services.projects import create_from_template

    tpl = Template.objects.filter(company=company).first()
    b = create_from_template(company, tpl, "翡翠湾 · 独栋", "翡翠湾 C 区 3 栋", datetime.date(2026, 10, 1),
                             datetime.date(2027, 9, 30), users["admin"])
    b_file = storage.save_bytes(b"%PDF-1.4 b-contract", "翡翠湾施工合同.pdf", users["admin"], b)
    routes = _project_routes()
    assert len(routes) > 80
    qc = client_for("qc")  # A 项目质检，不是 B 项目成员
    for method, pattern in routes:
        path = pattern.replace("<int:project_id>", str(b.id))
        path = re.sub(r"<int:\w+>", "1", path)
        path = re.sub(r"<str:\w+>", "2026-10-27" if "attendance" in path else "layout", path)
        r = getattr(qc, method)(path, {}, format="json")
        assert r.status_code == 403, (method, path, r.status_code)
        body = r.content.decode()
        assert r.json()["message"] == "无权访问该项目" and "翡翠湾" not in body, (method, path)
    # 跨项目文件直链
    r = qc.get(f"/api/files/{b_file.id}/content")
    assert r.status_code in (403, 404) and b"b-contract" not in r.content
    # 用自己项目的 URL 夹带别的项目对象 ID：按项目条件取对象，返回 404，不泄露数据
    from core.models import WorkItem
    b_wi = WorkItem.objects.filter(project=b).first()
    r = qc.get(f"/api/projects/{project.id}/work-items/{b_wi.id}")
    assert r.status_code == 404 and "翡翠湾" not in r.content.decode()


# ---------------------------------------------------------------- RULE-07 时区
from core.tests.test_check import check as _check, ev, make_cp  # noqa: E402,F401 — 复用 F06 测试的前置数据


@pytest.mark.django_db
def test_rule07_server_utc_display_beijing(project, client_for, ev, monkeypatch):
    import datetime
    import io
    import time
    import zipfile
    from unittest import mock

    from openpyxl import load_workbook

    monkeypatch.setenv("TZ", "UTC")  # 服务器系统时区设为 UTC
    time.tzset()
    try:
        utc_now = datetime.datetime(2026, 10, 27, 16, 30, tzinfo=datetime.timezone.utc)  # 北京时间 10-28 00:30
        qc = client_for("qc")
        with mock.patch("django.utils.timezone.now", return_value=utc_now):
            cp = make_cp(qc, project, ev["wi"], point="卫生间墙面线管走向")
            r = _check(qc, project, cp["id"], "一致", checkedAt="2026-01-01T08:00:00+08:00")  # 客户端时间应被忽略
            assert r.status_code == 201, r.content
        want = "2026-10-28T00:30:00+08:00"
        items = qc.get(f"/api/projects/{project.id}/work-items/{ev['wi'].id}/checkpoints").json()["items"]
        row = next(x for x in items if x["id"] == cp["id"])
        assert row["latestCheck"]["checkedAt"] == want
        det = qc.get(f"/api/projects/{project.id}/checkpoints/{cp['id']}").json()
        assert det["latestCheck"]["checkedAt"] == want
        # 证据包记录清单中的时间同样是北京时间
        owner = client_for("owner")
        job = owner.post(f"/api/projects/{project.id}/evidence-exports", {}, format="json").json()
        z = zipfile.ZipFile(io.BytesIO(b"".join(owner.get(f"/api/projects/{project.id}/evidence-exports/{job['id']}/download").streaming_content)))
        rows = list(load_workbook(io.BytesIO(z.read("记录清单.xlsx")))["记录清单"].iter_rows(values_only=True))
        times = [r[2] for r in rows[1:] if r[1] == "核对记录"]
        assert times and all(str(t).startswith("2026-10-28 00:30") for t in times), times
    finally:
        monkeypatch.delenv("TZ")
        time.tzset()
