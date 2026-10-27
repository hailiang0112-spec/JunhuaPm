"""plan 浏览器 E2E：FLOW-04 / 05 / 15 / 16 / 19 / 23（TDD 验收契约 FLOW 行「浏览器 E2E」部分；FLOW-23 契约只要求 API 集成测试，此处附带页面验证）。

前置：run.sh 已建库、seed_base、seed.py。业务操作都在真实页面完成；API 只用于读取断言所需事实。
"""
import base64
import csv
import io
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "_lib"))
from harness import WEB, Run  # noqa: E402

SEED = json.load(open(os.path.join(HERE, "seed.json")))
PID, P2, P3 = SEED["pid"], SEED["pidFlow04"], SEED["pidImport"]
BRICK = SEED["phases"]["砌墙"]
DL = os.path.join(HERE, "downloads")
os.makedirs(DL, exist_ok=True)
R = Run(HERE, expected_http=[r"/phases/\d+/plan", r"/attendance", r"/floors/\d+", r"/rooms/\d+", r"/work-items",
                             r"/plan/imports", r"/prototype-imports"])


def fetch(page, path):
    return page.evaluate("p => fetch('/api' + p).then(r => r.json())", path)


def r1(x):
    return float(f"{x + 1e-9:.1f}")


# ———————————————————————————————— FLOW-19（先改砌墙计划，FLOW-16 的 10-20~10-27 才在工期内） ————————————————————————————————
R.flow = "FLOW-19"


def f19():
    ctx, page = R.ctx("pm", 1440)
    page.goto(f"{WEB}/p/{PID}/phase/{BRICK}")
    page.wait_for_selector("text=工作项 · 完成情况")
    logs0 = len(fetch(page, f"/projects/{PID}/phases/{BRICK}/detail")["planLogs"]) if False else None
    page.click("button:has-text('编辑工段计划')")
    dates = page.locator("form#plan-form input[type=date]")
    R.check("砌墙原计划结束日 2026-10-18", dates.nth(1).input_value() == "2026-10-18", dates.nth(1).input_value())
    dates.nth(1).fill("2026-11-08")
    page.locator("form#plan-form textarea").last.fill("砌块到货延迟三周，书房 B 门洞待变更单签认")
    diff = page.locator("form#plan-form").inner_text()
    R.check("保存前预览结束日期前后值与工期影响 +21 天", "2026-10-18" in diff and "2026-11-08" in diff and "+21" in diff, diff[:200])
    page.click(".dlg-ft button:has-text('保存计划')")
    R.check("保存成功 toast", R.toast(page, "留档"), page.locator(".toast").all_inner_texts())
    page.wait_for_timeout(800)
    log = page.locator("#log-t").locator("xpath=ancestor::section[1]").inner_text()
    R.check("留档 1 条含时间、工段（砌墙工段页）、结束日期前后值、+21 天与原因",
            "砌墙" in page.locator("h1").first.inner_text() and re.search(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}", log)
            and all(x in log for x in ("计划结束日期", "2026-10-18", "2026-11-08", "+21", "砌块到货延迟三周")), log[:300])
    n1 = page.locator("#log-t").locator("xpath=ancestor::section[1]").locator("tbody tr").count()
    R.shot(page, "F10-03-planlog-1440")
    page.wait_for_timeout(3200)
    page.click("button:has-text('编辑工段计划')")
    R.check("原样打开表单提示「尚未修改任何字段…不会产生留档」", "不会产生留档" in page.locator("form#plan-form").inner_text())
    page.click(".dlg-ft button:has-text('保存计划')")
    page.wait_for_timeout(1500)
    n2 = page.locator("#log-t").locator("xpath=ancestor::section[1]").locator("tbody tr").count()
    R.check("原样保存不新增留档", n1 == n2 == 1, (n1, n2))
    if page.locator(".dialog").count():
        page.click(".dlg-ft button:has-text('取消')")
    page.click("button:has-text('编辑工段计划')")
    page.locator("form#plan-form input[type=date]").nth(0).fill("2026-11-20")
    page.click(".dlg-ft button:has-text('保存计划')")
    R.check("开始晚于结束被拒「开始日期不能晚于结束日期」", "开始日期不能晚于结束日期" in page.locator("form#plan-form").inner_text())
    page.click(".dlg-ft button:has-text('取消')")
    ctx.close()
    # 导出留档 CSV（审计页 / 总览台账）
    ctx, page = R.ctx("owner", 1440)
    page.goto(f"{WEB}/p/{PID}/audit")
    btn = page.locator("button:has-text('导出 CSV')").first
    btn.wait_for(timeout=10000)
    with page.expect_download() as d:
        btn.click()
    path = os.path.join(DL, "计划变更留档.csv")
    d.value.save_as(path)
    raw = open(path, "rb").read()
    text = raw.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    R.check("CSV 列为 记录时间、工段、变更项、变更前、变更后、工期影响（天）、变更原因备注",
            rows[0] == ["记录时间", "工段", "变更项", "变更前", "变更后", "工期影响（天）", "变更原因备注"], rows[0])
    R.check("CSV 含 1 条砌墙结束日期记录 +21", len(rows) == 2 and rows[1][1] == "砌墙" and rows[1][3] == "2026-10-18"
            and rows[1][4] == "2026-11-08" and rows[1][5].lstrip("+") == "21", rows[1:])
    ctx.close()


R.step("计划留档", f19)

# ———————————————————————————————— FLOW-16 ————————————————————————————————
R.flow = "FLOW-16"


def f16():
    ctx, page = R.ctx("pm", 1440)
    page.goto(f"{WEB}/p/{PID}/phase/{BRICK}")
    page.wait_for_selector("text=工作项 · 完成情况")
    page.click(".topbar button.primary:has-text('登记出勤')")
    page.fill("form#att-form input[type=date]", "2026-10-20")
    page.fill("form#att-form input[type=number]", "41")
    page.click(".dlg-ft button:has-text('保存')")
    R.check("41 人被拒「请填写 0–40 的整数」", "请填写 0–40 的整数" in page.locator("form#att-form").inner_text())
    page.fill("form#att-form input[type=number]", "6")
    page.click(".dlg-ft button:has-text('保存')")
    R.check("单日登记 10-20 6 人 toast", R.toast(page, "6"))
    page.wait_for_timeout(800)
    st = page.locator(".attstats").inner_text()
    R.check("单日保存后累计人工日 6、峰值 6 重算", re.search(r"累计人工日\s*6\b", st) and re.search(r"出勤峰值\s*6\b", st), st)
    page.wait_for_timeout(3200)
    page.click("button:has-text('批量登记出勤')")
    d = page.locator("form#batch-form input[type=date]")
    d.nth(0).fill("2026-10-21")
    d.nth(1).fill("2026-10-27")
    page.click("form#batch-form label:has-text('仅工作日（跳过周日）')")
    page.fill("form#batch-form input[type=number]", "5")
    page.wait_for_timeout(200)
    prev = page.locator("form#batch-form").inner_text()
    R.check("批量预览写入 6 天（跳过 10-25 周日）", re.search(r"将写入\s*6\s*天", prev), prev[-160:])
    page.click(".dlg-ft button:has-text('写入')")
    R.check("批量写入显示写入天数「已写入 6 天」", R.toast(page, "已写入 6 天"))
    page.wait_for_timeout(800)
    st = page.locator(".attstats").inner_text()
    R.check("批量后累计 36 人·日、日均 5.1、峰值 6、施工 7 天", all(re.search(p, st) for p in (r"累计人工日\s*36\b", r"日均出勤\s*5\.1\b", r"出勤峰值\s*6\b", r"施工天数\s*7\b")), st)
    R.shot(page, "F10-03-attendance-1440")
    ctx.close()
    for w in (1440, 375):
        ctx, page = R.ctx("owner", w)
        page.goto(f"{WEB}/p/{PID}")
        page.wait_for_selector("#cal-t")
        cal = page.locator("#cal-t").locator("xpath=ancestor::section[1]")
        R.check(f"{w} 日历为 2026 年 10 月", "2026 年 10 月" in cal.locator(".month").inner_text())
        cell = lambda d: cal.locator(f"td[title='2026-10-{d:02d}']")  # noqa: E731
        R.check(f"{w} 10-19 未登记（按 0 计、不推算）", cell(19).get_attribute("class").split()[0] == "cu" and "未登记" in cell(19).inner_text())
        R.check(f"{w} 10-21 为各工段之和 5+3=8、5–8 档", "8 人" in cell(21).inner_text() and "c2" in cell(21).get_attribute("class"))
        R.check(f"{w} 10-22 为 5+4=9、9–12 档", "9 人" in cell(22).inner_text() and "c3" in cell(22).get_attribute("class"))
        R.check(f"{w} 10-20 为 6 人、5–8 档", "6 人" in cell(20).inner_text() and "c2" in cell(20).get_attribute("class"))
        R.check(f"{w} 10-25（周日）未登记", "cu" in cell(25).get_attribute("class"))
        sm = cal.locator(".sumrow").inner_text()
        R.check(f"{w} 月合计 43 人·日、施工 7 天、峰值 9 人", "合计 43 人·日" in sm and "施工 7 天" in sm and "峰值 9 人" in sm, sm)
        R.check(f"{w} 图例五档", cal.locator(".legend span").count() == 6)
        if w == 375:
            R.check("375 总览无横向滚动（日历在 .scroll-x 内）", R.hs(page, "PAGE-F10-01 项目总览"))
        cal.screenshot(path=os.path.join(R.shots, f"F10-01-calendar-{w}.png"))
        ctx.close()


R.step("出勤", f16)

# ———————————————————————————————— FLOW-15 ————————————————————————————————
R.flow = "FLOW-15"
MODULES_F1001 = ["项目总览", "总体完成度", "计划基准", "工期进度", "今日出勤", "证据链四项指标", "六阶段时间轴", "楼层 × 阶段 完成度矩阵",
                 "短板 TOP 8", "房间问题聚焦 TOP 6", "每日出勤日历", "问题发现与闭环", "问题遗留清单", "计划变更台账", "各工段速览"]
MODULES_F1003 = ["工段进度", "计划基准", "计划开始 / 实际开始", "计划结束 / 实际结束", "工期 / 已施工天数", "出勤 · 峰值", "工作项 / 问题",
                 "计划变更留档", "工作项 · 完成情况", "本工段楼层 · 房间分布", "每日出勤人数", "出勤明细台账", "工作完成验收图片", "发现问题与闭环"]


def expected_overall(page):
    tot = 0.0
    for name, pid in SEED["phases"].items():
        items = [w for w in fetch(page, f"/projects/{PID}/work-items?phase={pid}")["items"] if not w["isDisabled"]]
        sw = sum(w["weight"] for w in items)
        pp = sum(w["weight"] * w["progress"] for w in items) / sw if sw else 0
        tot += SEED["weights"][name] * pp
    return r1(tot / sum(SEED["weights"].values()))


def f15():
    for w in (1440, 375):
        ctx, page = R.ctx("owner", w)
        page.goto(f"{WEB}/p/{PID}")
        page.wait_for_selector("text=证据链四项指标")
        page.wait_for_timeout(800)
        body = page.inner_text("body")
        if w == 1440:
            exp = expected_overall(page)
            m = re.search(r"总体完成度\s*([\d.]+)\s*%", body)
            R.check(f"总体完成度按加权公式 = {exp}%（1 位小数）", m and float(m.group(1)) == exp, m.group(0) if m else body[:200])
            m2 = re.search(r"计划基准 ([\d.]+)% · (领先|滞后) ([\d.]+) 个百分点", body)
            ok = m2 and m and r1(abs(float(m.group(1)) - float(m2.group(1)))) == float(m2.group(3)) and \
                (m2.group(2) == "领先") == (float(m.group(1)) >= float(m2.group(1)))
            R.check("显示计划基准与领先/滞后百分点（= 完成度 − 计划基准）", ok, m2.group(0) if m2 else "")
            R.check("周期 2026-09-19 ~ 2027-08-30、已用 39 天 / 总 346 天", "2026-09-19 ~ 2027-08-30" in body and "已用 39 天 / 总 346 天" in body)
            cards = page.locator(".evi .metric")
            got = {cards.nth(i).inner_text().split("\n")[0].strip(): cards.nth(i).locator(".vl").inner_text().replace("\n", "") for i in range(cards.count())}
            R.check("四项指标 70.0% / 75.0% / 90.0% / 75.0%", [got.get(k) for k in ("覆盖前发现比例", "隐蔽检查点合规率", "最新签认版使用率", "变更先签后做比例")]
                    == ["70.0%", "75.0%", "90.0%", "75.0%"], got)
            tg = page.locator(".evi").inner_text()
            R.check("目标线 90% / 100% / 100% / 100%", tg.count("目标 90%") == 1 and tg.count("目标 100%") == 3)
            missing = [x for x in MODULES_F1001 if x not in body]
            R.check("原型全部总览模块均在", not missing, missing)
            page.click(".chiprow button:has-text('铺砖')")
            page.wait_for_timeout(1200)
            na = page.locator(".evi .metric .vl.na").count()
            evi = page.locator(".evi").inner_text()
            R.check("无数据工段四项均显示「暂无数据」而非 0% 或 100%", na == 4 and "0.0%" not in evi and "100.0%" not in evi, (na, evi[:200]))
            R.shot(page, "F10-01-no-data-1440")
            page.click(".chiprow button:has-text('全项目')")
        R.check(f"{w} 总览无横向滚动", R.hs(page, "PAGE-F10-01 项目总览"))
        R.shot(page, f"F10-01-overview-{w}")
        page.goto(f"{WEB}/p/{PID}/floors")
        page.wait_for_selector("text=房间 × 阶段 完成度矩阵")
        R.check(f"{w} 楼层 / 房间页无横向滚动", R.hs(page, "PAGE-F10-02 楼层房间"))
        R.shot(page, f"F10-02-floors-{w}")
        if w == 1440:
            sec = page.locator("#rm-t").locator("xpath=ancestor::section[1]")
            sec.scroll_into_view_if_needed()
            cell = sec.locator("a[aria-label^='书房 B 砌墙']").first
            R.check("房间矩阵有「书房 B × 砌墙」单元格", cell.count() == 1)
            cell.click()
            page.wait_for_selector(".noteline", timeout=10000)
            note = page.locator(".noteline").inner_text()
            R.check("矩阵单元格跳到砌墙工段详情且筛选为该房间", f"/phase/{BRICK}" in page.url and "书房 B" in note and "二楼" in note, (page.url, note))
            body = page.inner_text("body")
            missing = [x for x in MODULES_F1003 if x not in body]
            R.check("工段详情包含原型全部模块", not missing, missing)
            R.shot(page, "F10-03-studyB-1440")
        else:
            page.goto(f"{WEB}/p/{PID}/phase/{BRICK}")
            page.wait_for_selector("text=工作项 · 完成情况")
            R.check("375 工段详情无横向滚动", R.hs(page, "PAGE-F10-03 工段详情"))
            R.shot(page, "F10-03-phase-375")
        ctx.close()


R.step("总览", f15)

# ———————————————————————————————— FLOW-04 ————————————————————————————————
R.flow = "FLOW-04"


def f04():
    ctx, page = R.ctx("pm", 1440)
    page.goto(f"{WEB}/p/{PID}/structure")
    page.wait_for_selector("#fl-t")
    # 新增并改名房间（二楼）
    page.locator("#fl-t").locator("xpath=ancestor::section[1]").locator("tbody tr", has_text="二楼").first.click()
    page.click("button:has-text('新增房间')")
    page.fill("form#name-form input", "储物间")
    page.select_option("form#name-form select", label="二楼")
    page.click(".dlg-ft button:has-text('保存')")
    R.check("新增房间 toast（矩阵与筛选项已同步）", R.toast(page, "已新增"))
    page.wait_for_timeout(3300)
    rooms = page.locator("#rm-t").locator("xpath=ancestor::section[1]")
    rooms.locator("tr", has_text="储物间").locator("button:has-text('改名')").click()
    page.fill("form#name-form input", "二楼储物间")
    page.click(".dlg-ft button:has-text('保存')")
    R.check("房间改名 toast", R.toast(page, "已改名"))
    page.wait_for_timeout(800)
    names = [r["name"] for r in fetch(page, f"/projects/{PID}/structure")["rooms"]]
    R.check("改名即时反映在结构数据中", "二楼储物间" in names and "储物间" not in names)
    page.goto(f"{WEB}/p/{PID}/issues")
    page.wait_for_timeout(1000)
    opts = fetch(page, f"/projects/{PID}/issues/options")
    R.check("改名即时反映在筛选项中（问题筛选房间）", any(r["name"] == "二楼储物间" for f in opts["floors"] for r in f["rooms"]))
    # 删除无记录楼层
    page.goto(f"{WEB}/p/{PID}/structure")
    page.wait_for_selector("#fl-t")
    page.click("button:has-text('新增楼层')")
    page.fill("form#name-form input", "阁楼")
    page.click(".dlg-ft button:has-text('保存')")
    R.toast(page, "已新增")
    page.wait_for_timeout(3300)
    fl = page.locator("#fl-t").locator("xpath=ancestor::section[1]")
    fl.locator("tr", has_text="阁楼").locator("button:has-text('删除')").click()
    page.click(".dlg-ft button:has-text('确认删除')")
    R.check("删除无记录楼层 toast", R.toast(page, "删除"))
    page.wait_for_timeout(800)
    R.check("楼层已删除", not any(f["name"] == "阁楼" for f in fetch(page, f"/projects/{PID}/structure")["floors"]))
    # 删除有记录的房间（书房 B）→ 只能停用
    page.wait_for_timeout(2600)
    fl.locator("tbody tr", has_text="二楼").first.click()
    rooms.locator("tr", has_text="书房 B").locator("button:has-text('删除')").click()
    page.wait_for_timeout(300)
    if page.locator(".dlg-ft button:has-text('确认删除')").count():
        page.click(".dlg-ft button:has-text('确认删除')")
    page.wait_for_selector(".blocked", timeout=8000)
    R.check("有记录房间删除被拒「已有记录，只能停用」", "已有记录，只能停用" in page.locator(".dialog").inner_text())
    R.shot(page, "F03-02-delete-blocked-1440")
    page.click(".dlg-ft button:has-text('确认停用')")
    R.check("停用 toast「已停用，历史记录保持可查」", R.toast(page, "已停用，历史记录保持可查"))
    page.wait_for_timeout(800)
    # 停用后不再出现在新增选项；房间与楼层不匹配被拒
    page.click(".topbar button:has-text('新增工作项')")
    form = page.locator("form#item-form")
    form.locator("select").nth(1).select_option(label="二楼")
    room_opts = form.locator("select").nth(2).locator("option").all_inner_texts()
    R.check("停用后书房 B 不再出现在新增选项", "书房 B" not in room_opts and "二楼储物间" in room_opts, room_opts)
    form.locator("input").first.fill("二楼卫生间 A 隔墙及管道井")
    form.locator("select").nth(2).select_option(label="卫生间 A")
    form.locator("select").nth(1).select_option(label="一楼")
    form.locator("input[type=number]").nth(0).fill("2")
    form.locator("input[type=number]").nth(1).fill("0")
    page.click(".dlg-ft button:has-text('保存工作项')")
    page.wait_for_timeout(800)
    R.check("楼层房间不匹配被拒「房间与楼层不匹配」", "房间与楼层不匹配" in form.inner_text() + " ".join(page.locator(".toast").all_inner_texts()), form.inner_text()[-200:])
    form.locator("select").nth(1).select_option(label="二楼")
    form.locator("select").nth(2).select_option(label="卫生间 A")
    form.locator("textarea").nth(1).fill("管道井封堵密实，闭水 24 小时无渗漏（项目经理填写）")
    page.click(".dlg-ft button:has-text('保存工作项')")
    R.check("新增工作项 toast", R.toast(page, "二楼卫生间 A 隔墙及管道井"))
    page.wait_for_timeout(800)
    item = max((w for w in fetch(page, f"/projects/{PID}/work-items")["items"] if w["name"] == "二楼卫生间 A 隔墙及管道井"), key=lambda w: w["id"])
    R.check("未填工作说明按规则库首个命中规则（管道井）补齐", "管道井" in (item["description"] or ""), item["description"])
    R.check("已填验收标准不被覆盖", item["acceptanceCriteria"] == "管道井封堵密实，闭水 24 小时无渗漏（项目经理填写）", item["acceptanceCriteria"])
    page.goto(f"{WEB}/p/{PID}/work-items/{SEED['studyB']}")
    page.wait_for_timeout(1500)
    R.check("停用房间的历史工作项仍可查", "书房 B" in page.inner_text("body"))
    ctx.close()
    # 砌墙进度 60.0%（项目 2：权重 2/2/1、完成度 100/50/0）
    for w in (1440, 375):
        ctx, page = R.ctx("pm", w)
        page.goto(f"{WEB}/p/{P2}/structure")
        page.click("[role=tab]:has-text('工段')")
        page.wait_for_selector("#ph-t")
        row = page.locator("#ph-t").locator("xpath=ancestor::section[1]").locator("tr", has_text="砌墙").first
        txt = row.inner_text()
        R.check(f"{w} 砌墙进度 (2×100+2×50+1×0)÷5 = 60.0%，状态「进行中」", "60.0%" in txt and "进行中" in txt, txt)
        R.check(f"{w} 工段进度与状态不可手改（行内无输入框）", row.locator("input, select").count() == 0)
        R.check(f"{w} 项目结构页无横向滚动", R.hs(page, "PAGE-F03-02 项目结构"))
        R.shot(page, f"F03-02-phases-{w}")
        ctx.close()




# ———————————————————————————————— FLOW-05 ————————————————————————————————
R.flow = "FLOW-05"
SHEETS = ["施工阶段", "工作事项", "补充事项", "填写说明", "取值"]


def check_xlsx(path, label):
    wb = openpyxl.load_workbook(path)
    R.check(f"{label} 含 5 张表", wb.sheetnames == SHEETS, wb.sheetnames)
    ws = wb["工作事项"]
    hdr = [c.value for c in ws[1]]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in r)]
    if rows:
        wi, pi = hdr.index("权重(1-5)"), hdr.index("完成度(%)")
        di = hdr.index("计划完成日期")
        R.check(f"{label} 数值列为数字", all(isinstance(r[wi], (int, float)) and isinstance(r[pi], (int, float)) for r in rows))
        R.check(f"{label} 日期列为文本", all(r[di] is None or isinstance(r[di], str) for r in rows))
    return wb


def f05():
    ctx, page = R.ctx("pm", 1440)
    page.goto(f"{WEB}/p/{PID}/import")
    page.wait_for_selector("text=Excel 计划模板与导出")
    with page.expect_download() as d:
        page.click("button:has-text('下载空白模板')")
    p1 = os.path.join(DL, "空白模板.xlsx")
    d.value.save_as(p1)
    check_xlsx(p1, "空白模板")
    with page.expect_download() as d:
        page.click("button:has-text('导出当前计划')")
    p2 = os.path.join(DL, "当前计划.xlsx")
    d.value.save_as(p2)
    wb = check_xlsx(p2, "导出计划")
    ws = wb["工作事项"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    # 错误文件：阶段名错误、楼层房间不匹配、日期无法识别、权重 7
    bad = openpyxl.load_workbook(p2)
    b = bad["工作事项"]
    b.cell(3, col["施工阶段"], "抹灰")
    b.cell(4, col["所属楼层"], "一楼")
    b.cell(4, col["所属房间 / 区域"], "卫生间 A")
    b.cell(5, col["计划完成日期"], "十月底")
    b.cell(6, col["权重(1-5)"], 7)
    bp = os.path.join(DL, "错误计划.xlsx")
    bad.save(bp)
    page.set_input_files(".drop input[type=file]", files=[{"name": "错误计划.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": open(bp, "rb").read()}])
    page.wait_for_selector("text=处错误", timeout=20000)
    errs = page.locator(".errtbl tbody tr").all_inner_texts()
    joined = " | ".join(errs)
    R.check("错误文件逐条列出 4 处错误", len(errs) >= 4, joined[:300])
    R.check("错误逐条含工作表与行号（第 3/4/5/6 行：阶段名、楼层房间、日期、权重）",
            all(any(f"第 {n} 行" in e and "工作事项" in e and k in e for e in errs) for n, k in ((3, "阶段名"), (4, "不匹配"), (5, "无法识别"), (6, "权重"))), joined[:400])
    R.check("错误文件不允许导入（无执行导入按钮）", page.locator("button:has-text('执行导入')").count() == 0)
    R.shot(page, "F03-03-errors-1440")
    # 合规文件：新增 1 项、改未签认项说明、改已签认项说明
    good = openpyxl.load_workbook(p2)
    g = good["工作事项"]
    names = {g.cell(r, col["工作事项名称"]).value: r for r in range(2, g.max_row + 1)}
    signed_row = names[SEED["signedItemName"]]
    g.cell(signed_row, col["工作说明"], "（导入试图修改）改用 200mm 砌块")
    draft = [w for w in fetch(page, f"/projects/{PID}/work-items?phase={SEED['phases']['水电']}")["items"] if w["baselineStatus"] == "草稿"]
    unsigned_name = next(w["name"] for w in draft if w["name"] in names)
    g.cell(names[unsigned_name], col["工作说明"], "导入更新：线管改用阻燃 PVC 管，接线盒处留 150mm 余量，穿线后做绝缘测试；顶部预留发泡剂封堵")
    new_row = g.max_row + 1
    src = names[unsigned_name]
    for c in range(1, len(hdr) + 1):
        g.cell(new_row, c, g.cell(src, c).value)
    g.cell(new_row, col["序号"], new_row - 1)
    g.cell(new_row, col["工作事项名称"], "负一楼保姆间新风机电源预留")
    g.cell(new_row, col["完成度(%)"], 0)
    gp = os.path.join(DL, "修正计划.xlsx")
    good.save(gp)
    before = next(w for w in fetch(page, f"/projects/{PID}/work-items")["items"] if w["id"] == SEED["signedItem"])
    page.set_input_files(".drop input[type=file]", files=[{"name": "修正计划.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": open(gp, "rb").read()}])
    page.wait_for_selector("button:has-text('执行导入')", timeout=20000)
    page.click("button:has-text('执行导入')")
    summ = page.locator(".dialog").inner_text()
    R.check("执行前摘要列出新增、更新与已签认跳过", "新增" in summ and "已签认" in summ, summ.replace("\n", " ")[:200])
    page.click(".dlg-ft button:has-text('确认导入')")
    R.check("导入完成 toast", R.toast(page, "导入"), page.locator(".toast").all_inner_texts())
    page.wait_for_timeout(1000)
    res = page.inner_text("body")
    R.check("结果逐项列出「已签认，需走变更单」", SEED["signedItemName"] in res and "已签认，需走变更单" in res)
    R.shot(page, "F03-03-result-1440")
    items = fetch(page, f"/projects/{PID}/work-items")["items"]
    after = next(w for w in items if w["id"] == SEED["signedItem"])
    R.check("已签认项字段完全不变", all(after[k] == before[k] for k in ("description", "acceptanceCriteria", "cost", "plannedFinish", "weight", "drawingVersionId")))
    R.check("未签认项按文件更新", any(w["name"] == unsigned_name and "发泡剂" in (w["description"] or "") for w in items))
    R.check("新增项已导入", any(w["name"] == "负一楼保姆间新风机电源预留" for w in items))
    ctx.close()
    ctx, page = R.ctx("pm", 375)
    page.goto(f"{WEB}/p/{PID}/import")
    page.wait_for_selector("text=Excel 计划模板与导出")
    R.check("375 数据导入导出页无横向滚动", R.hs(page, "PAGE-F03-03 导入导出"))
    R.shot(page, "F03-03-375")
    ctx.close()



R.flow = "FLOW-05"
R.step("Excel 导入导出", f05)  # 先于 FLOW-04：FLOW-04 会停用书房 B、新增同名工作项，改变导出的计划
R.flow = "FLOW-04"
R.step("项目结构", f04)

# ———————————————————————————————— FLOW-23（附带：契约为 API 集成测试） ————————————————————————————————
R.flow = "FLOW-23"
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def backup():
    png = "data:image/png;base64," + base64.b64encode(PNG).decode()
    return {
        "phase": {"p1": {"start": "2026-09-19", "end": "2026-10-18", "crew": "瓦工班 · 陈师傅（6人）", "w": 0.08}},
        "tasks": {"p1": [
            {"n": "放线定位、材料进场验收", "note": "", "std": "", "plan": "09-19", "act": "09-19", "owner": "陈师傅", "fl": "all", "rm": "all",
             "w": 1, "pct": 100, "files": ["F1"]},
            {"n": "负一楼卫生间新砌隔墙及管道井封堵", "plan": "09-22", "owner": "陈师傅", "fl": "b1", "rm": "b1-wc", "w": 2, "pct": 100, "files": ["F2"]}]},
        "att": {"p1": {"2026-09-19": 8, "2026-09-20": 6}},
        "issues": [{"id": "N123456", "pid": "p1", "t": "负一楼车库门洞偏位 30mm", "d": "2026-09-25", "lv": "重要", "fix": "拆除重砌",
                    "owner": "陈师傅", "due": "2026-10-02", "fl": "b1", "rm": "b1-car", "st": "未闭环"}],
        "edit": {}, "delIssues": {}, "closed": {"N123456": "2026-10-01"},
        "photos": {"p1": [{"cap": "负一楼卫生间管道井封堵", "st": "已验收", "src": png, "date": "2026-09-23"}]},
        "upto": "", "log": [{"t": "2026-09-30 18:20", "pid": "p1", "pname": "砌墙", "note": "材料晚到",
                             "items": [{"f": "end", "from": "2026-10-15", "to": "2026-10-18", "dd": 3}]}],
        "files": [{"id": "F1", "sec": "task", "fo": "", "name": "砌块出厂检测报告-MU10.pdf", "type": "application/pdf", "size": 204800},
                  {"id": "F2", "sec": "task", "fo": "", "name": "负一楼卫生间管道井封堵-完成面.jpg", "type": "image/jpeg", "size": 1024},
                  {"id": "F3", "sec": "ct", "fo": "c1", "name": "开工交底记录-20260919.pdf", "type": "application/pdf", "size": 4096}],
        "folders": {"ct": [{"id": "c1", "name": "开工资料"}]},
        "extra": [{"id": "E1", "title": "业主要求增加车库储物架", "desc": "口头约定", "date": "2026-09-28", "due": "2026-10-10", "st": "处理中",
                   "owner": "张工", "files": []}],
        "notice": [{"id": "M1", "board": "owner", "who": "刘先生", "text": "周六上午来看现场", "ts": "2026-10-08 09:00", "re": ""},
                   {"id": "M2", "board": "owner", "who": "张工", "text": "好的，九点在门口等", "ts": "2026-10-08 09:30", "re": "M1"}],
        "noticeName": {},
    }


def f23():
    ctx, page = R.ctx("admin", 1440)
    page.goto(f"{WEB}/p/{P3}/import")
    page.wait_for_selector("text=原型 v1.20 备份导入")
    inp = page.locator("input[type=file][accept*='json']").first
    inp.set_input_files(files=[{"name": "君华别墅项目数据-2026-10-09.json", "mimeType": "application/json",
                                "buffer": json.dumps(backup(), ensure_ascii=False).encode("utf-8")}])
    page.wait_for_timeout(1500)
    for label in ("开始导入", "确认导入", "执行导入"):
        b = page.locator(f".dialog button:has-text('{label}'), button:has-text('{label}')").last
        if b.count() and b.is_visible():
            b.click()
            page.wait_for_timeout(800)
    page.wait_for_timeout(1500)
    body = page.inner_text("body")
    R.shot(page, "F03-03-prototype-1440")
    R.check("导入结果出现（已导入）", "已导入" in body or R.toast(page, "导入"), body[-300:])
    items = fetch(page, f"/projects/{P3}/work-items?phase={fetch(page, f'/projects/{P3}/structure')['phases'][0]['id']}")["items"]
    R.check("工作事项逐类导入（砌墙整段替换为 2 项）", len([w for w in items if not w["isDisabled"]]) == 2, len(items))
    R.check("只有元数据的资料标记「文件缺失，需重新上传」", "文件缺失，需重新上传" in body or "文件缺失" in body)
    snap = len(fetch(page, f"/projects/{P3}/work-items")["items"])
    inp.set_input_files(files=[{"name": "君华别墅备份.json", "mimeType": "application/json",
                                "buffer": json.dumps({"name": "君华别墅备份", "rows": []}).encode()}])
    page.wait_for_timeout(1500)
    for label in ("开始导入", "确认导入"):
        b = page.locator(f".dialog button:has-text('{label}')").last
        if b.count() and b.is_visible():
            b.click()
            page.wait_for_timeout(800)
    msg = page.inner_text("body") + " ".join(page.locator(".toast").all_inner_texts())
    R.check("非原型 JSON 被拒「不是有效的原型备份文件」", "不是有效的原型备份文件" in msg)
    R.check("被拒后项目数据不变", len(fetch(page, f"/projects/{P3}/work-items")["items"]) == snap)
    ctx.close()


R.step("原型导入", f23)
R.finish()
